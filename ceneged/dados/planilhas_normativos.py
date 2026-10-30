"""Exportação da planilha dos normativos e modelos para download.

Parte do programa Levantamento de Materiais (Ceneged)."""

from ceneged.comum.texto import fmt_num
from ceneged.dados.planilhas import salvar_planilha_composicoes
from ceneged.dados.normativos import (CHAVE_ALTURAS, CHAVE_POSTES, ESTRUTURAS_NORMA,
    TABELAS_NORMA, _COLS_ALTURAS, _COLS_POSTES, _COLS_TABELA)

# ---------------------------------------------------------------------------
#  Planilha dos normativos a partir do que está no sistema (para editar no Excel)
# ---------------------------------------------------------------------------
def _txt_faixa(v):
    if v is None or v == "":
        return None
    if isinstance(v, list):
        if v and isinstance(v[0], list):
            return "; ".join("/".join(fmt_num(x) for x in par) for par in v)
        return f"{fmt_num(v[0])} a {fmt_num(v[1])}" if len(v) == 2 else ", ".join(fmt_num(x) for x in v)
    return v


def salvar_planilha_normativos(caminho):
    """Grava estruturas, tabelas, postes e alturas do sistema numa planilha no mesmo modelo
    da planilha de carregar (dá para editar no Excel e carregar de novo)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    azul = PatternFill("solid", fgColor="1F4E78")
    amarelo = PatternFill("solid", fgColor="FFF2CC")

    def cab(ws, cols, larg=None):
        ws.append(cols)
        for c, nome in enumerate(cols, 1):
            x = ws.cell(1, c)
            x.font, x.fill = Font(bold=True, color="FFFFFF"), azul
            x.alignment = Alignment(wrap_text=True, vertical="center")
            ws.column_dimensions[get_column_letter(c)].width = (larg or {}).get(nome, 13)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}1"

    ws = wb.active
    ws.title = "Leia-me"
    for t in ["PLANILHA DOS NORMATIVOS (gerada pelo programa, com as alterações feitas nele)",
              "Para carregar: tela Materiais extras > Normativos > \"Carregar planilha dos normativos\".",
              "Cada aba carregada substitui o que estava no sistema. Não mude os títulos das colunas.",
              "Faixa de bitola: \"35 a 50\". Pares: \"25/25; 35/35\". Tipo = cunha marca conector cunha.",
              "Postes: Topo/Base = diâmetro (circular) ou face A (DT); face B nas colunas seguintes."]:
        ws.append([t])
    ws["A1"].font = Font(bold=True, size=13, color="1F4E78")
    ws.column_dimensions["A"].width = 120

    est = ESTRUTURAS_NORMA.conteudo
    if est:
        cols = ["Normativo", "Nº", "Estrutura", "Pág.", "Seção / Tensão", "Item", "Cód. Nordeste", "Descrição",
                "Und", "Qtd DT", "Qtd DT 36,2 kV", "Qtd Circular", "Qtd (única)", "Variável", "A definir?",
                "Informação necessária para definir o material", "Referência no normativo", "Observação"]
        ws = wb.create_sheet("Estruturas")
        cab(ws, cols, {"Estrutura": 16, "Seção / Tensão": 24, "Descrição": 44, "Referência no normativo": 24,
                       "Informação necessária para definir o material": 50})
        num = {}
        for e in est:
            num[e["norma"]] = num.get(e["norma"], 0) + 1
            for i in e["itens"]:
                ws.append([e["norma"], num[e["norma"]], e["nome"], e.get("pag"), i.get("secao"), i.get("item"),
                           i.get("ne") or ("A DEFINIR" if i.get("adef") else ""), i.get("desc"), i.get("und"),
                           i.get("qdt"), i.get("qdt36"), i.get("qcirc"), i.get("qunica"), i.get("variavel"),
                           "SIM" if i.get("adef") else None, i.get("info"), i.get("ref"), i.get("obs")])
                ws.cell(ws.max_row, 7).number_format = "@"
                if i.get("adef"):
                    for c in range(1, len(cols) + 1):
                        ws.cell(ws.max_row, c).fill = amarelo
    tabs = TABELAS_NORMA.conteudo
    normais = {k: v for k, v in tabs.items() if not k.startswith("DIS-ETE-011|")}
    if normais:
        cols = [n for n, ch in _COLS_TABELA if ch not in ("se", "br")]
        ws = wb.create_sheet("Tabelas")
        cab(ws, cols, {"Título": 32, "Informação necessária": 44, "Descrição": 44})
        for k, t in normais.items():
            norma, _, tab = k.partition("|")
            for l in t.get("linhas", []):
                nua = t.get("modo") == "nua"
                val = {"ne": l.get("ne"), "descricao": l.get("descricao"), "b": l.get("b"),
                       "p": _txt_faixa(l.get("p")), "d": _txt_faixa(l.get("d")), "f": l.get("f"), "n": l.get("n"),
                       "pares": _txt_faixa(l.get("pares")), "mat": l.get("mat"),
                       "tipo": "cunha" if l.get("cunha") is True else l.get("tipo"), "kv": l.get("kv"),
                       "comp": None if nua else l.get("comp"),
                       "cunha": None if l.get("cunha") is True else l.get("cunha"),
                       "comp_cod": l.get("comp") if nua else None,
                       "postes": "; ".join(l.get("postes") or []) or None,
                       "padrao": "SIM" if l.get("padrao") else None}
                ws.append([norma, tab, t.get("titulo"), t.get("modo"), t.get("info")] +
                          [val.get(ch) for n, ch in _COLS_TABELA if ch not in (None, "se", "br")])
    if CHAVE_POSTES in tabs:
        ws = wb.create_sheet("Postes")
        cab(ws, _COLS_POSTES, {"Descrição": 30, "Uso": 30})
        for l in tabs[CHAVE_POSTES].get("linhas", []):
            dt = l.get("tipo") == "dt"
            ws.append(["DT" if dt else "Circular", l.get("l"), l.get("rn"), l.get("rn_a"), l.get("ne"),
                       l.get("descricao"), l.get("a") if dt else l.get("topo"), l.get("A") if dt else l.get("base"),
                       l.get("b") if dt else None, l.get("B") if dt else None, l.get("uso")])
    if CHAVE_ALTURAS in tabs:
        ws = wb.create_sheet("Altura de aplicação")
        cab(ws, _COLS_ALTURAS, {"Situação": 26, "Observação": 60})
        for l in tabs[CHAVE_ALTURAS].get("linhas", []):
            ws.append([l.get("norma"), l.get("situacao"), l.get("d"), l.get("obs")])
    wb.save(caminho)


# ---------------------------------------------------------------------------
#  Modelos de planilha (botões "Baixar modelo" na tela Materiais extras)
# ---------------------------------------------------------------------------
def _aba_leia(wb, titulo, linhas, primeira=False):
    from openpyxl.styles import Font
    ws = wb.active if primeira else wb.create_sheet("Como preencher")
    if primeira:
        ws.title = "Como preencher"
    ws.append([titulo])
    ws["A1"].font = Font(bold=True, size=13, color="1F4E78")
    for t in linhas:
        ws.append([t])
        if t and t == t.upper() and len(t) > 3:
            ws.cell(ws.max_row, 1).font = Font(bold=True)
    ws.column_dimensions["A"].width = 130
    return ws


def salvar_modelo_composicoes(caminho):
    """Modelo da planilha de materiais extras (materiais interligados), com exemplo."""
    exemplo = [
        {"id": "RES-0190", "categoria": "EXEMPLO - BAIXA TENSÃO", "nome": "FLBIT", "apelidos": ["FL-BIT"],
         "auto": True, "itens": [
            {"nome": "FIO COBRE 750 V 2,50 MM2 PT", "codigo": "2221001", "und": "M", "qtd": 0.5, "mov": "Aplicação"},
            {"nome": "ISOLADOR ROLDANA PORC 76,0MM 1350DAN.", "codigo": "2300000", "und": "CDA", "qtd": 1,
             "mov": "Aplicação"},
            {"nome": "ARMACAO SECUN ACO CARB 1 ESTR.150X 110MM", "codigo": "3417025", "und": "CDA", "qtd": 1,
             "mov": "Aplicação"}]},
        {"id": "RES-0167", "categoria": "EXEMPLO - BAIXA TENSÃO", "nome": "IT-R", "apelidos": [], "auto": True,
         "itens": [
            {"nome": "CONETOR PERF 16,0- 95,0/ 4,0- 35,0", "codigo": "2412001", "und": "CDA", "qtd": 4,
             "mov": "Aplicação"},
            {"nome": "PARAFUSO CAB QUAD ACO 16X 200", "codigo": "3480305", "und": "CDA", "qtd": 1,
             "mov": "Aplicação"}]}]
    salvar_planilha_composicoes(caminho, exemplo)
    from openpyxl import load_workbook
    wb = load_workbook(caminho)
    _aba_leia(wb, "MODELO – PLANILHA DE MATERIAIS EXTRAS (materiais interligados)", [
        "Para carregar: tela Materiais extras > \"Carregar planilha\". Ela SUBSTITUI todos os materiais principais.",
        "A PRIMEIRA ABA é a que o programa lê (a Planilha1). Esta aba de explicação pode ficar ou ser apagada.",
        "",
        "COLUNAS (linha de títulos; a ordem pode mudar, os nomes não)",
        "• Item / Material – nome da classificação, do material principal ou do material que vem junto.",
        "• Cód. Material – código do material (só nas linhas de material que vem junto).",
        "• UND – CDA, M, KG, CJ...",
        "• Qtd por unid. lançada – quantidade para cada unidade do principal (cabo: por metro).",
        "• Movimento – Aplicação ou Retirada.",
        "• Nomes no projeto – outros nomes do principal nos balões, separados por vírgula (opcional).",
        "• Automático – Sim/Não: se o principal entra sozinho quando aparece no balão (opcional).",
        "",
        "TIPOS DE LINHA (pela cor da célula do nome, ou pela posição)",
        "• VERDE = classificação (ex.: BAIXA TENSÃO). Vale para os principais abaixo dela.",
        "• AZUL CLARO = material principal. Escreva \"CÓDIGO  ·  NOME\" (ex.: RES-0190  ·  FLBIT) ou só o nome.",
        "• BRANCO = materiais que vêm junto (com código e quantidade), logo abaixo do principal.",
        "",
        "EXEMPLO",
        "• As linhas da aba Planilha1 são só exemplo: a classificação começa com \"EXEMPLO\" e o programa não",
        "  carrega nada que esteja numa classificação \"EXEMPLO\". Apague e coloque os seus materiais.",
        "• O jeito mais fácil: use \"Baixar planilha (com as alterações)\", que já sai neste modelo com tudo do sistema.",
    ])
    wb.save(caminho)


def salvar_modelo_normativos(caminho):
    """Modelo da planilha dos normativos: explicação + abas com os títulos certos e exemplos."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    azul = PatternFill("solid", fgColor="1F4E78")
    cinza = PatternFill("solid", fgColor="EDEDED")
    _aba_leia(wb, "MODELO – PLANILHA DOS NORMATIVOS", [
        "Para carregar: tela Materiais extras > Normativos > \"Carregar planilha dos normativos\".",
        "Cada aba carregada SUBSTITUI a mesma parte no sistema. Pode mandar só uma aba (ex.: só Postes).",
        "O programa acha cada aba pelos TÍTULOS DAS COLUNAS: nome da aba, ordem das colunas e colunas a mais",
        "não importam; o texto dos títulos tem que ser igual ao deste modelo.",
        "Linhas cinza com Normativo começando com \"EXEMPLO\" são ignoradas: apague e coloque os dados de verdade.",
        "O jeito mais fácil: na janela das tabelas, \"Baixar planilha dos normativos (com as alterações)\" já sai",
        "neste modelo com tudo do sistema.",
        "",
        "ABA ESTRUTURAS – uma linha por material de cada estrutura",
        "• Normativo (DIS-NOR-014...), Estrutura (nome como vem no balão: N1, CE3, FLBIT), Item (letra do desenho).",
        "• Cód. Nordeste: código; vazio ou \"A DEFINIR\" quando o material sai de uma tabela.",
        "• Qtd DT / Qtd DT 36,2 kV / Qtd Circular / Qtd (única): quantidade conforme poste e tensão; \"-\" = não usa.",
        "• Seção / Tensão: \"Material Específico para 15 kV\" / \"36,2 kV\", \"Rede Primária\" / \"Rede Secundária\".",
        "• A definir? = SIM e Referência no normativo = \"Tabela 3 do DIS-NOR-014\": o material sai da aba Tabelas.",
        "• Material com \"(POSTE 9 M)\" na descrição: entra só no poste dessa altura.",
        "",
        "ABA TABELAS – uma linha por opção de material",
        "• Normativo + Tabela iguais aos da \"Referência no normativo\" (ex.: DIS-NOR-014 + Tabela 3).",
        "• Escolha por: poste, bitola, cabo, fase, neutro, mensageiro, multiplexado, principal, derivacao, estribo,",
        "  terminacao, pares, nua, nua_bitola, laco, unico, manual.",
        "• Bitolas em mm² (4 AWG = 21,2; 2 = 33,6; 1/0 = 53,5; 4/0 = 107,2; 336,4 MCM = 170,5).",
        "• Faixa: \"35 a 50\". Pares principal/derivação: \"25/25; 35/35\". Tipo = cunha: conector cunha.",
        "",
        "ABA POSTES – dimensões da DIS-ETE-011",
        "• Tipo (DT ou Circular), Comprimento (m), Resistência (daN) como no balão (400/12 = 400 e 12).",
        "• Topo/Base (mm): diâmetro do circular ou face A do DT; Topo/Base face B só no DT.",
        "• Cód. Nordeste vazio = poste que não existe no Nordeste (não é usado para descobrir o tipo).",
        "",
        "ABA ALTURA DE APLICAÇÃO – distância do topo onde a peça é presa",
        "• Situação: MT; BT sozinha no poste; BT com MT no poste; Suporte de equipamento (Normativo = Equipamento).",
    ], primeira=True)

    def aba(nome, cols, exemplos, larg=None):
        ws = wb.create_sheet(nome)
        ws.append(cols)
        for c, n in enumerate(cols, 1):
            x = ws.cell(1, c)
            x.font, x.fill = Font(bold=True, color="FFFFFF"), azul
            x.alignment = Alignment(wrap_text=True, vertical="center")
            ws.column_dimensions[get_column_letter(c)].width = (larg or {}).get(n, 13)
        for r in exemplos:
            ws.append(r)
            for c in range(1, len(cols) + 1):
                ws.cell(ws.max_row, c).fill = cinza
        ws.freeze_panes = "A2"

    aba("Estruturas", ["Normativo", "Nº", "Estrutura", "Pág.", "Seção / Tensão", "Item", "Cód. Nordeste",
                       "Descrição", "Und", "Qtd DT", "Qtd DT 36,2 kV", "Qtd Circular", "Qtd (única)", "Variável",
                       "A definir?", "Informação necessária para definir o material", "Referência no normativo",
                       "Observação"],
        [["EXEMPLO DIS-NOR-014", 7, "FLBIT", 21, None, "fb", "3417025", "ARMACAO SECUN ACO CARB 1 ESTR", "CDA",
          1, None, 1, None, "-", None, None, None, None],
         ["EXEMPLO DIS-NOR-014", 7, "FLBIT", 21, None, "ff", "A DEFINIR", "CINTA DE AÇO CARBONO", "CDA",
          "-", None, 1, None, "POSTE", "SIM", "Diâmetro do poste no ponto de fixação", "Tabela 3 do DIS-NOR-014",
          None]],
        {"Descrição": 40, "Informação necessária para definir o material": 40, "Referência no normativo": 24})
    aba("Tabelas", [n for n, ch in _COLS_TABELA if ch not in ("se", "br")],
        [["EXEMPLO DIS-NOR-014", "Tabela 3", "Cintas de Aço", "poste", "diâmetro do poste no ponto de fixação",
          "3416055", "CINTA POSTE CIRC ACO CARB 160 MM"],
         ["EXEMPLO DIS-NOR-014", "Tabela 5", "Conector Perfurante 1 Derivação", "pares", "",
          "2412001", "CONECTOR PERF 16,0-95,0 / 4,0-35,0", None, None, None, None, None, "25/25; 35/35; 70/35"]],
        {"Título": 30, "Informação necessária": 36, "Descrição": 40, "Pares principal/derivação (mm²)": 24})
    aba("Postes", _COLS_POSTES,
        [["DT", 12, 400, 200, "3300153", "EXEMPLO POSTE CONCRETO DT 12M 400DAN", 140, 476, 110, 350, "Rede primária"],
         ["Circular", 12, 1000, None, "3302058", "EXEMPLO POSTE CONCRETO CIRCULAR 12M 1000DAN", 230, 470, None, None,
          "Rede primária"]], {"Descrição": 34, "Uso": 20})
    aba("Altura de aplicação", _COLS_ALTURAS,
        [["EXEMPLO DIS-NOR-018", "MT", 0.2, "cruzeta a 200 mm do topo"],
         ["EXEMPLO DIS-NOR-014", "BT com MT no poste", 1.5, "BT abaixo da MT"]],
        {"Normativo": 20, "Situação": 24, "Observação": 40})
    wb.save(caminho)

"""Leitura e gravação das planilhas (materiais, constantes, carteira, composições).

Parte do programa Levantamento de Materiais (Ceneged)."""

import os
import re

from ceneged.comum.texto import (_sem_acento, chave_nome, normaliza_codigo)
from ceneged.dados.composicoes import (MOVIMENTOS, _kit_partes, _num, titulo_kit)
from ceneged.dados.equipes import (CAMPOS_OBRA, CORES_STATUS, OPC_SN, OPC_SNP, STATUS_OBRA,
    _opcao, data_br)

def ler_planilha_materiais(caminho):
    """Lê a planilha de cadastro: colunas "Código" e "Material" (títulos na 1ª linha).
    -> [(código, nome)]. Sem a coluna Código, os códigos ficam vazios."""
    ext = os.path.splitext(caminho)[1].lower()
    t_nome = ("MATERIAIS", "MATERIAL", "NOME", "DESCRICAO", "DESCRICAOSUCINTA", "NOMEDOMATERIAL")
    t_cod = ("CODIGO", "COD", "CODIGOMATERIAL", "CODMATERIAL", "CODIGODOMATERIAL", "CODIGONEOENERGIA")

    def acha(linhas):
        for r, lin in enumerate(linhas[:30]):
            ks = [chave_nome(v) if v is not None else "" for v in lin]
            cn = next((c for c, k in enumerate(ks) if k in t_nome), None)
            if cn is not None:
                cc = next((c for c, k in enumerate(ks) if k in t_cod), None)
                return r, cn, cc
        return None
    linhas, pos = [], None
    if ext in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        wb = load_workbook(caminho, read_only=True, data_only=True)
        for ws in wb.worksheets:
            linhas = [list(r) for r in ws.iter_rows(values_only=True)]
            pos = acha(linhas)
            if pos:
                break
        wb.close()
    elif ext in (".csv", ".txt"):
        bruto = open(caminho, "rb").read()
        for cod in ("utf-8-sig", "cp1252", "latin-1"):
            try:
                texto = bruto.decode(cod)
                break
            except UnicodeDecodeError:
                continue
        # Excel em português separa por ";". Sem ";" nem tab, é uma coluna só
        # (a vírgula fica no nome: 3xS04-13,8kV).
        sep = ";" if ";" in texto else ("\t" if "\t" in texto else None)
        linhas = [(ln.split(sep) if sep else [ln]) for ln in texto.splitlines()]
        linhas = [[c.strip().strip('"') for c in ln] for ln in linhas]
        pos = acha(linhas)
    elif ext == ".xls":
        raise ValueError("Planilha no formato antigo (.xls).\nAbra no Excel e use "
                         "\"Salvar como\" > Pasta de Trabalho do Excel (.xlsx).")
    else:
        raise ValueError("Use uma planilha do Excel (.xlsx) ou um arquivo .csv.")
    if not pos:
        raise ValueError("Não encontrei os títulos da planilha.\n\n"
                         "Na primeira linha: célula A1 = Código e célula B1 = Material.\n"
                         "Da linha 2 para baixo: o código e o nome de cada material.")
    r, cn, cc = pos
    itens = []

    def texto_de(v):
        if isinstance(v, float) and v.is_integer():
            v = int(v)
        return re.sub(r"\s+", " ", str(v)).strip() if v is not None else ""
    for lin in linhas[r + 1:]:
        nome = texto_de(lin[cn] if cn < len(lin) else None)
        cod = texto_de(lin[cc] if cc is not None and cc < len(lin) else None)
        if nome:
            itens.append((normaliza_codigo(cod) if cod else "", nome))
    return itens


def salvar_modelo_planilha(caminho):
    """Planilha modelo para o cadastro (Código | Material + instruções)."""
    if caminho.lower().endswith(".xlsx"):
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        wb = Workbook()
        ws = wb.active
        ws.title = "Materiais"
        ws["A1"], ws["B1"] = "Código", "Material"
        for c in ("A1", "B1"):
            ws[c].font = Font(bold=True, color="FFFFFF")
            ws[c].fill = PatternFill("solid", fgColor="014FA0")
        ws.column_dimensions["A"].width = 14
        ws.column_dimensions["B"].width = 55
        ins = wb.create_sheet("Como preencher")
        ins.column_dimensions["A"].width = 95
        linhas = ["COMO PREENCHER A PLANILHA DE CADASTRO DE MATERIAIS",
                  "",
                  "1. Use só a aba \"Materiais\".",
                  "2. Célula A1: Código     Célula B1: Material",
                  "3. Da linha 2 para baixo: o código e o nome de cada material, um por linha,",
                  "   com o nome exatamente como a empresa usa.",
                  "4. Linhas vazias e nomes repetidos são ignorados.",
                  "5. Salve como .xlsx e use o botão \"Carregar planilha\" no programa.",
                  "",
                  "Exemplo:",
                  "   A2: 3417025    B2: ARMACAO SECUN ACO CARB 1 ESTR",
                  "   A3: 2300000    B3: ISOLADOR ROLDANA PORCELANA",
                  "   A4: 3480310    B4: PARAFUSO M16X250"]
        for i, t in enumerate(linhas, 1):
            ins.cell(i, 1, t).alignment = Alignment(wrap_text=True)
        ins["A1"].font = Font(bold=True, color="014FA0", size=13)
        wb.save(caminho)
    else:
        with open(caminho, "w", newline="", encoding="utf-8-sig") as f:
            f.write("Código;Material\n")
UNIDADES = ("CDA", "UN", "M", "KG", "CJ", "L")


def ler_planilha_constantes(caminho):
    """Planilha com uma coluna de constante (número) e uma de descrição; código opcional."""
    from openpyxl import load_workbook
    wb = load_workbook(caminho, data_only=True, read_only=True)
    linhas = [list(r) for r in wb.active.iter_rows(values_only=True)]
    wb.close()
    cab = None
    for i, lin in enumerate(linhas[:10]):
        ks = [chave_nome(v) if v is not None else "" for v in lin]
        cc = next((j for j, k in enumerate(ks) if k.startswith("CONST")), None)
        cd = next((j for j, k in enumerate(ks) if k.startswith("DESCRI") or k.startswith("MATERIA")), None)
        if cc is not None and cd is not None:
            cab = (i, cc, cd, next((j for j, k in enumerate(ks) if k.startswith("COD")), None))
            break
    if not cab:
        raise ValueError("Não encontrei os títulos: uma coluna CONSTANTE e uma DESCRIÇÃO (código é opcional).")
    i0, cc, cd, co = cab
    itens = []
    for lin in linhas[i0 + 1:]:
        d = lin[cd] if cd < len(lin) else None
        k = _num(lin[cc]) if cc < len(lin) else None
        if not d or k is None:
            continue
        cod = lin[co] if co is not None and co < len(lin) else ""
        itens.append({"codigo": normaliza_codigo(cod or ""), "descricao": re.sub(r"\s+", " ", str(d)).strip(),
                      "constante": k})
    return itens


def _campo_da_coluna(titulo):
    k = chave_nome(titulo)
    for prefixo, campo in (("DATADODESLIG", "data_desl"), ("DATADESLIG", "data_desl"), ("DATACONC", "conclusao"),
                           ("DATAINI", "inicio"), ("DATAFIM", "fim"), ("DATAFINAL", "fim"),
                           ("PROJETO", "projeto"), ("NOTA", "nota"), ("SUPERVISOR", "supervisor"),
                           ("ENCARREGADO", "encarregado"), ("VIABILIZ", "viabilizada"), ("REQUISI", "requisicao"),
                           ("EQUIPE", "equipe"), ("DESLIG", "desligamento"), ("STATUS", "status"),
                           ("OBS", "obs")):
        if k.startswith(prefixo):
            return campo
    return None


def ler_planilha_carteira(caminho):
    """Planilha da carteira (títulos como PROJETO, NOTA, SUPERVISOR...) -> [dados]."""
    from openpyxl import load_workbook
    wb = load_workbook(caminho, data_only=True, read_only=True)
    for ws in wb.worksheets:
        linhas = list(ws.iter_rows(values_only=True))
        for i, lin in enumerate(linhas[:20]):
            cols = {}
            for j, v in enumerate(lin):
                c = _campo_da_coluna(v) if v is not None else None
                if c and c not in cols:
                    cols[c] = j
            if "projeto" in cols and len(cols) >= 4:
                saida = []
                for lin2 in linhas[i + 1:]:
                    d = {}
                    for c, j in cols.items():
                        v = lin2[j] if j < len(lin2) else None
                        if c in ("inicio", "fim", "conclusao"):
                            try:
                                v = data_br(v)
                            except ValueError:
                                v = str(v).strip()
                        elif hasattr(v, "strftime"):
                            v = v.strftime("%d/%m/%Y")
                        elif isinstance(v, float) and v.is_integer():
                            v = str(int(v))
                        d[c] = "" if v is None else str(v).strip()
                    if not d.get("projeto"):
                        continue
                    for c in ("viabilizada", "requisicao"):
                        d[c] = _opcao(d.get(c), OPC_SNP)
                    d["desligamento"] = _opcao(d.get("desligamento"), OPC_SN)
                    d["status"] = _opcao(d.get("status"), STATUS_OBRA)
                    for c in ("supervisor", "encarregado", "equipe"):
                        d[c] = re.sub(r"\s+", " ", d.get(c, "")).strip().upper()
                    saida.append(d)
                wb.close()
                return saida
    wb.close()
    raise ValueError("Não encontrei os títulos da carteira (PROJETO, NOTA, SUPERVISOR, EQUIPE...).")


def salvar_planilha_carteira(caminho, obras):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook()
    ws = wb.active
    ws.title = "Carteira de obra"
    ws.append([t for _, t, _ in CAMPOS_OBRA])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="014FA0")
        c.alignment = Alignment(horizontal="center", vertical="center")
    col_status = [k for k, _, _ in CAMPOS_OBRA].index("status") + 1
    for o in obras:
        lin = []
        for k, _, _ in CAMPOS_OBRA:
            v = o.get(k, "")
            lin.append(int(v) if k in ("projeto", "nota") and str(v).isdigit() else v)
        ws.append(lin)
        cor = CORES_STATUS.get(o.get("status"))
        if cor:
            cel = ws.cell(ws.max_row, col_status)
            cel.fill = PatternFill("solid", fgColor=cor[1].lstrip("#"))
            cel.font = Font(color="FFFFFF", bold=True)
    for i, (_, t, w) in enumerate(CAMPOS_OBRA, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = max(len(t) + 4, w / 7)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(caminho)


def _cor_tipo(c):
    """Cor de fundo da célula -> 'categoria' (verde), 'principal' (azul claro), 'cabecalho' (escuro) ou None."""
    try:
        f = c.fill
        if not f or f.fill_type in (None, "none"):
            return None
        cor = f.fgColor
        if cor.type != "rgb" or not isinstance(cor.rgb, str):
            return "?"
        h = cor.rgb[-6:]
        r, g, b = (int(h[i:i + 2], 16) for i in (0, 2, 4))
    except Exception:
        return None
    if min(r, g, b) >= 245:
        return None
    if g > r + 40 and g > b + 25:
        return "categoria"
    if max(r, g, b) < 130:
        return "cabecalho"
    if b >= r and min(r, g, b) >= 150:
        return "principal"
    return "?"


def ler_planilha_composicoes(caminho):
    """Planilha 'materiais interligados':
         Material | Código | UND | Qtd por unid. lançada | Movimento  [| Nomes no projeto | Automático]
       verde = classificação, azul claro = material principal, branco = materiais que vêm junto.
    -> (kits, avisos)"""
    from openpyxl import load_workbook
    wb = load_workbook(caminho, data_only=True)
    try:
        wbf = load_workbook(caminho, data_only=False)
    except Exception:
        wbf = None
    ws = wb.active
    wsf = wbf[ws.title] if wbf else None
    linhas = list(ws.iter_rows())
    # cabeçalho
    col = {"nome": 0, "codigo": 1, "und": 2, "qtd": 3, "mov": 4, "apelidos": None, "auto": None}
    ini = 0
    for i, row in enumerate(linhas[:15]):
        txts = [_sem_acento(str(c.value or "")).upper() for c in row]
        if any("MATERIAL" in t or "ITEM" in t for t in txts) and any("COD" in t or "QTD" in t for t in txts):
            for j, t in enumerate(txts):
                if ("ITEM" in t or "MATERIAL" in t) and "COD" not in t and "QTD" not in t and col["nome"] == 0 and j < 3:
                    col["nome"] = j
                if "COD" in t:
                    col["codigo"] = j
                elif t.startswith("UN"):
                    col["und"] = j
                elif "QTD" in t or "QUANT" in t:
                    col["qtd"] = j
                elif "MOVIM" in t:
                    col["mov"] = j
                elif "NOMES" in t or "APELIDO" in t:
                    col["apelidos"] = j
                elif "AUTOM" in t:
                    col["auto"] = j
            ini = i + 1
            break

    def val(row, chave):
        j = col[chave]
        return row[j].value if j is not None and j < len(row) else None

    avisos = []
    sem_valor = 0
    tipos = []
    for i, row in enumerate(linhas[ini:], ini):
        nome = re.sub(r"\s+", " ", str(val(row, "nome") or "")).strip()
        if not nome:
            tipos.append(None)
            continue
        cod, q = val(row, "codigo"), val(row, "qtd")
        if wsf is not None and col["codigo"] is not None:
            fc = wsf.cell(i + 1, col["codigo"] + 1).value
            if cod is None and isinstance(fc, str) and fc.startswith("="):
                sem_valor += 1
        if (cod not in (None, "") and _num(cod) is not None) or _num(q) is not None:
            tipos.append("item")
        else:
            tipos.append(_cor_tipo(row[col["nome"]]) or "?")
    # sem cor reconhecida: título seguido de outro título = classificação
    for i, t in enumerate(tipos):
        if t in ("?", "cabecalho"):
            prox = next((x for x in tipos[i + 1:] if x), None)
            tipos[i] = "categoria" if prox in ("categoria", "principal", "?", "cabecalho", None) else "principal"
    kits, cat, atual, soltos = [], "SEM CLASSIFICAÇÃO", None, 0
    for row, t in zip(linhas[ini:], tipos):
        if not t:
            continue
        nome = re.sub(r"\s+", " ", str(val(row, "nome") or "")).strip()
        if t == "categoria":
            cat, atual = nome, None
            if _sem_acento(nome).upper().startswith("EXEMPLO"):
                cat = None                    # classificação de exemplo do modelo: não carrega
        elif cat is None:
            continue
        elif t == "principal":
            kid, kn = _kit_partes(nome)
            atual = {"id": kid, "categoria": cat, "nome": kn, "itens": []}
            ap = val(row, "apelidos")
            if ap:
                atual["apelidos"] = [x.strip() for x in str(ap).split(",") if x.strip()]
            au = val(row, "auto")
            if au not in (None, ""):
                atual["auto"] = not _sem_acento(str(au)).upper().startswith("N")
                atual["auto_planilha"] = True
            kits.append(atual)
        elif atual is None:
            soltos += 1
        else:
            atual["itens"].append({"nome": nome, "codigo": val(row, "codigo") or "", "und": val(row, "und") or "CDA",
                                   "qtd": val(row, "qtd"), "mov": val(row, "mov") or MOVIMENTOS[0]})
    if sem_valor:
        avisos.append(f"{sem_valor} código(s) estão em fórmulas sem valor salvo (ligadas a outra planilha). "
                      "Abra a planilha no Excel, salve e carregue de novo.")
    if soltos:
        avisos.append(f"{soltos} material(is) apareceram antes de qualquer material principal e foram ignorados.")
    if not kits:
        avisos.append("Nenhum material principal encontrado (linhas em azul claro).")
    return kits, avisos


def salvar_planilha_composicoes(caminho, kits):
    """Mesma estrutura da planilha de entrada (pode ser carregada de volta)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook()
    ws = wb.active
    ws.title = "Planilha1"
    ws.append(["Item / Material", "Cód. Material", "UND", "Qtd por unid. lançada", "Movimento",
               "Nomes no projeto", "Automático"])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E78")
        c.alignment = Alignment(horizontal="center")
    verde = PatternFill("solid", fgColor="00A000")
    azul = PatternFill("solid", fgColor="DDEBF7")
    cat = None
    for k in kits:
        if k["categoria"] != cat:
            cat = k["categoria"]
            ws.append([cat])
            for c in ws[ws.max_row][:5]:
                c.fill = verde
            ws.cell(ws.max_row, 1).font = Font(bold=True, color="FFFFFF")
        ws.append([titulo_kit(k), None, None, None, None, ", ".join(k["apelidos"]),
                   "Sim" if k["auto"] else "Não"])
        for c in ws[ws.max_row][:5]:
            c.fill = azul
        ws.cell(ws.max_row, 1).font = Font(bold=True, color="1F4E78")
        for i in k["itens"]:
            cod = i["codigo"]
            q = i["qtd"]
            ws.append(["    " + i["nome"], int(cod) if cod.isdigit() else (cod or None), i["und"],
                       int(q) if q == int(q) else q, i["mov"]])
    for c, w in zip("ABCDEFG", (52, 14, 7, 20, 20, 28, 11)):
        ws.column_dimensions[c].width = w
    ws.freeze_panes = "A2"
    wb.save(caminho)

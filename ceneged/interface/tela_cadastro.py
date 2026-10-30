"""Tela do cadastro de materiais.

Parte do programa Levantamento de Materiais (Ceneged)."""

import re
import tkinter as tk
from tkinter import filedialog
from tkinter import messagebox
from tkinter import ttk

from ceneged.comum.texto import chave_nome
from ceneged.dados.cadastro import CADASTRO
from ceneged.dados.planilhas import (ler_planilha_materiais, salvar_modelo_planilha)
from ceneged.interface.tema import (AZUL, BORDA, BRANCO, CINZA, FONTE, FUNDO, LARANJA, TEXTO,
    VERDE, VERM_ESC, amostra, cartao, secao, tabela)
from ceneged.interface.dialogos import EscolherNome

class TelaCadastro(tk.Frame):
    """Cadastro dos nomes reais dos materiais (planilha) e conferência do de-para."""
    COMO = {"igual": ("Igual ao cadastro", "igual"), "parecido": ("Parecido ({:.0%})", "parecido"),
            "manual": ("Escolhido à mão", "manual"), "nao": ("Não encontrado", "nao"),
            "": ("Sem cadastro", "vazio")}

    # ---- permissão: quem não pode editar só consulta
    def _ed(self, w):
        self.__dict__.setdefault("edicao", []).append(w)
        return w

    def aplicar_permissao(self, pode):
        for w in getattr(self, "edicao", []):
            w.state(["!disabled"] if pode else ["disabled"])
        if not pode and not hasattr(self, "_aviso_leitura"):
            self._aviso_leitura = tk.Label(self, text="Você pode consultar, mas não alterar. "
                                                      "Fale com o administrador.",
                                           bg="#FFF4D6", fg="#7A5600", font=(FONTE, 9), anchor="w", padx=10, pady=6)
            self._aviso_leitura.pack(fill="x", padx=20, pady=(0, 8), after=self.winfo_children()[0])

    def __init__(self, pai, app):
        super().__init__(pai, bg=FUNDO)
        self.app = app
        cab = tk.Frame(self, bg=FUNDO)
        cab.pack(fill="x", padx=20, pady=(16, 10))
        ttk.Label(cab, text="Cadastro de materiais", style="Titulo.TLabel").pack(anchor="w")
        ttk.Label(cab, style="Sub.TLabel", wraplength=1100, justify="left",
                  text="Os nomes que vêm no projeto nem sempre são os que a empresa usa. O programa procura "
                       "cada nome do projeto no cadastro e, no relatório, usa o nome verdadeiro do cadastro.").pack(anchor="w")

        corpo = tk.Frame(self, bg=FUNDO)
        corpo.pack(fill="both", expand=True, padx=20, pady=(0, 16))
        corpo.columnconfigure(0, weight=2, uniform="c")
        corpo.columnconfigure(1, weight=3, uniform="c")
        corpo.rowconfigure(0, weight=1)

        # ---- esquerda: cadastro + explicação da planilha
        esq = tk.Frame(corpo, bg=FUNDO)
        esq.grid(row=0, column=0, sticky="nsew", padx=(0, 10))
        c1 = cartao(esq)
        c1.pack(fill="both", expand=True)
        self.sec_cad = secao(c1, "Materiais cadastrados")
        self.sec_cad.pack(anchor="w", padx=14, pady=(12, 6))
        bts = tk.Frame(c1, bg=BRANCO)
        bts.pack(fill="x", padx=14)
        self._ed(ttk.Button(bts, text="Carregar planilha", style="Azul.TButton", command=self.carregar)).pack(side="left")
        ttk.Button(bts, text="Baixar modelo", command=self.baixar_modelo).pack(side="left", padx=6)
        ttk.Button(bts, text="Baixar cadastro", command=self.baixar_cadastro).pack(side="left")
        self._ed(ttk.Button(bts, text="Limpar", command=self.limpar)).pack(side="right")
        busca = tk.Frame(c1, bg=BRANCO)
        busca.pack(fill="x", padx=14, pady=(10, 6))
        tk.Label(busca, text="Buscar:", bg=BRANCO, fg=CINZA, font=(FONTE, 9)).pack(side="left")
        self.v_busca = tk.StringVar()
        self.v_busca.trace_add("write", lambda *a: self.mostrar_cadastro())
        ttk.Entry(busca, textvariable=self.v_busca).pack(side="left", fill="x", expand=True, padx=(6, 0))
        # edição de um material
        ed = tk.Frame(c1, bg=BRANCO)
        ed.pack(side="bottom", fill="x", padx=14, pady=(6, 12))
        tk.Label(ed, text="Código", bg=BRANCO, fg=CINZA, font=(FONTE, 8)).grid(row=0, column=0, sticky="w")
        tk.Label(ed, text="Material (nome da empresa)", bg=BRANCO, fg=CINZA, font=(FONTE, 8)).grid(
            row=0, column=1, sticky="w", padx=(6, 0))
        self.v_cod, self.v_nome = tk.StringVar(), tk.StringVar()
        self._ed(ttk.Entry(ed, textvariable=self.v_cod, width=12)).grid(row=1, column=0, sticky="we")
        self._ed(ttk.Entry(ed, textvariable=self.v_nome)).grid(row=1, column=1, sticky="we", padx=(6, 0))
        ed.columnconfigure(1, weight=1)
        be = tk.Frame(ed, bg=BRANCO)
        be.grid(row=2, column=0, columnspan=2, sticky="we", pady=(6, 0))
        self._ed(ttk.Button(be, text="Adicionar novo", style="Azul.TButton",
                            command=self.mat_adicionar)).pack(side="left")
        self._ed(ttk.Button(be, text="Salvar alteração", command=self.mat_salvar)).pack(side="left", padx=4)
        self._ed(ttk.Button(be, text="Excluir", command=self.mat_excluir)).pack(side="left")
        ttk.Button(be, text="Limpar campos", command=self.mat_limpar).pack(side="right")
        self.lbl_mat = tk.Label(ed, text="Clique num material da lista para editar, ou preencha e clique em "
                                         "\"Adicionar novo\".", bg=BRANCO, fg=CINZA, font=(FONTE, 8), anchor="w")
        self.lbl_mat.grid(row=3, column=0, columnspan=2, sticky="we", pady=(4, 0))
        self.mat_sel = None
        self.tree_cad = tabela(c1, [("cod", "Código", 90), ("mat", "Material (nome da empresa)", 300)], altura=4)
        self.tree_cad.quadro.pack(fill="both", expand=True, padx=14)
        self.tree_cad.bind("<<TreeviewSelect>>", lambda e: self.mat_abrir())
        self.tree_cad.bind("<Delete>", lambda e: self.mat_excluir())

        c2 = cartao(esq)
        c2.pack(fill="x", pady=(12, 0))
        secao(c2, "Como deve ser a planilha").pack(anchor="w", padx=14, pady=(12, 8))
        linha = tk.Frame(c2, bg=BRANCO)
        linha.pack(fill="x", padx=14, pady=(0, 14))
        self._mini_planilha(linha).pack(side="left", anchor="n")
        passos = ("1.  Excel (.xlsx) ou .csv, uma aba só.\n"
                  "2.  Na linha 1, os títulos:  A1 = Código  e  B1 = Material\n"
                  "3.  Da linha 2 para baixo, um material por linha: o código e o nome "
                  "exatamente como a empresa usa.\n"
                  "4.  Linhas vazias e nomes repetidos são ignorados.\n"
                  "5.  Ao carregar, você escolhe se a planilha substitui o cadastro "
                  "ou é acrescentada a ele.\n\n"
                  "Use \"Baixar modelo\" para receber a planilha já pronta.")
        tk.Label(linha, text=passos, bg=BRANCO, fg=TEXTO, font=(FONTE, 9), justify="left",
                 wraplength=300, anchor="nw").pack(side="left", fill="both", expand=True, padx=(14, 0))

        # ---- direita: nomes do projeto -> nome do cadastro
        dir_ = cartao(corpo)
        dir_.grid(row=0, column=1, sticky="nsew")
        secao(dir_, "Nome no projeto  >  nome no relatório").pack(anchor="w", padx=14, pady=(12, 2))
        tk.Label(dir_, bg=BRANCO, fg=CINZA, font=(FONTE, 9), justify="left", anchor="w", wraplength=700,
                 text="Nomes encontrados nos PDFs abertos e o nome que vai para o relatório. Se algum estiver "
                      "errado, dê dois cliques na linha e escolha o nome certo; a escolha fica guardada para "
                      "os próximos projetos.").pack(fill="x", padx=14, pady=(0, 8))
        self.tree_map = tabela(dir_, [("proj", "Nome no projeto", 150), ("cod", "Código", 80),
                                      ("cad", "Nome no relatório", 280), ("como", "Situação", 130)], altura=12)
        self.tree_map.quadro.pack(fill="both", expand=True, padx=14)
        self.tree_map.tag_configure("igual", foreground=VERDE)
        self.tree_map.tag_configure("parecido", foreground=LARANJA)
        self.tree_map.tag_configure("manual", foreground=AZUL)
        self.tree_map.tag_configure("nao", foreground=VERM_ESC)
        self.tree_map.tag_configure("vazio", foreground=CINZA)
        self.tree_map.bind("<Double-1>", lambda e: self.escolher())
        bts2 = tk.Frame(dir_, bg=BRANCO)
        bts2.pack(fill="x", padx=14, pady=12)
        self._ed(ttk.Button(bts2, text="Escolher nome…", style="Azul.TButton", command=self.escolher)).pack(side="left")
        self._ed(ttk.Button(bts2, text="Voltar ao automático", command=self.automatico)).pack(side="left", padx=6)
        legenda = tk.Frame(bts2, bg=BRANCO)
        legenda.pack(side="right")
        for txt, cor in (("igual", VERDE), ("parecido", LARANJA), ("à mão", AZUL), ("não encontrado", VERM_ESC)):
            amostra(legenda, cor, txt)

    def _mini_planilha(self, pai):
        """Desenho de como a planilha deve ficar."""
        g = tk.Frame(pai, bg=BORDA)
        cel = dict(font=(FONTE, 8), bd=0, padx=6, pady=3, anchor="w")
        tk.Label(g, text="", bg="#EEF1F6", width=2, **cel).grid(row=0, column=0, padx=(1, 1), pady=(1, 1), sticky="nsew")
        for c, (letra, larg) in enumerate((("A", 8), ("B", 24)), 1):
            tk.Label(g, text=letra, bg="#EEF1F6", fg=CINZA, width=larg, **{**cel, "anchor": "center"}).grid(
                row=0, column=c, padx=(0, 1), pady=(1, 1), sticky="nsew")
        linhas = [(("Código", "Material"), True), (("3417025", "ARMACAO SECUN ACO CARB 1 ESTR"), False),
                  (("2300000", "ISOLADOR ROLDANA PORCELANA"), False), (("3480310", "PARAFUSO M16X250"), False),
                  (("…", "…"), False)]
        for i, (ts, tit) in enumerate(linhas, 1):
            tk.Label(g, text=str(i), bg="#EEF1F6", fg=CINZA, width=2, **{**cel, "anchor": "center"}).grid(
                row=i, column=0, padx=(1, 1), pady=(0, 1), sticky="nsew")
            for c, t in enumerate(ts, 1):
                tk.Label(g, text=t, bg=AZUL if tit else BRANCO, fg=BRANCO if tit else TEXTO,
                         **{**cel, "font": (FONTE, 8, "bold" if tit else "normal")}).grid(
                    row=i, column=c, padx=(0, 1), pady=(0, 1), sticky="nsew")
        return g

    # ---- dados
    def atualizar(self):
        self.mostrar_cadastro()
        self.mostrar_nomes()

    def mostrar_cadastro(self):
        t = self.tree_cad
        t.delete(*t.get_children())
        k = chave_nome(self.v_busca.get())
        for i, m in enumerate(CADASTRO.materiais):
            cod = CADASTRO.codigo_de(m)
            if not k or k in chave_nome(m) or k in cod:
                t.insert("", "end", iid=f"m{i}", values=(cod, m))
        if self.mat_sel and t.exists(self.mat_sel):
            t.selection_set(self.mat_sel)
            t.see(self.mat_sel)
        self.sec_cad.rotulo.config(text=f"Materiais cadastrados ({len(CADASTRO.materiais)})")

    def mostrar_nomes(self):
        t = self.tree_map
        sel = t.selection()
        t.delete(*t.get_children())
        nomes = self.app.nomes_dos_projetos()
        if not nomes:
            t.insert("", "end", values=("", "", "Abra e processe os PDFs na Página inicial.", ""), tags=("vazio",))
            return
        for n in sorted(nomes, key=str.upper):
            real, como, nota = CADASTRO.resolver(n)
            txt, tag = self.COMO.get(como, ("", "nao"))
            cod = CADASTRO.codigo_de(real)
            if como in ("igual", "parecido", "manual") and CADASTRO.codigos and not cod:
                txt, tag = txt + " (sem código)", "nao"
            t.insert("", "end", iid="n:" + n, values=(n, cod, real, txt.format(nota)), tags=(tag,))
        for s in sel:
            if t.exists(s):
                t.selection_set(s)
                t.see(s)

    def carregar(self):
        if not self.app.pode_editar:
            return
        cam = filedialog.askopenfilename(title="Planilha de materiais",
                                         filetypes=[("Planilhas", "*.xlsx *.xlsm *.csv"), ("Todos", "*.*")])
        if not cam:
            return
        try:
            nomes = ler_planilha_materiais(cam)
        except Exception as ex:
            messagebox.showerror("Planilha", str(ex))
            return
        if not nomes:
            messagebox.showwarning("Planilha", "Achei os títulos, mas não há materiais embaixo deles.")
            return
        sem_cod = sum(1 for c, n in nomes if not c)
        substituir = True
        if CADASTRO.materiais:
            r = messagebox.askyesnocancel(
                "Cadastro", f"A planilha tem {len(nomes)} material(is).\n\n"
                            f"Sim  = substituir o cadastro atual ({len(CADASTRO.materiais)} materiais)\n"
                            f"Não  = acrescentar aos que já estão cadastrados")
            if r is None:
                return
            substituir = r
        CADASTRO.definir_materiais(nomes, substituir)
        self.app.cadastro_mudou()
        messagebox.showinfo("Cadastro", f"Cadastro atualizado: {len(CADASTRO.materiais)} material(is)." +
                            (f"\n\nAtenção: {sem_cod} linha(s) sem código." if sem_cod else ""))

    def baixar_modelo(self):
        try:
            import openpyxl  # noqa: F401
            tipos, ext = [("Excel", "*.xlsx")], ".xlsx"
        except Exception:
            tipos, ext = [("CSV", "*.csv")], ".csv"
        cam = filedialog.asksaveasfilename(defaultextension=ext, filetypes=tipos,
                                           initialfile="modelo_cadastro_materiais" + ext)
        if cam:
            try:
                salvar_modelo_planilha(cam)
            except Exception as ex:
                messagebox.showerror("Modelo", str(ex))
                return
            messagebox.showinfo("Modelo", f"Modelo salvo:\n{cam}")

    def limpar(self):
        if not self.app.pode_editar:
            return
        if CADASTRO.materiais and messagebox.askyesno(
                "Cadastro", "Apagar todos os materiais cadastrados?\n(As regras de materiais extras continuam.)"):
            CADASTRO.definir_materiais([], substituir=True)
            self.app.cadastro_mudou()

    # ---- edição do cadastro
    def mat_abrir(self):
        s = self.tree_cad.selection()
        if not s:
            return
        self.mat_sel = s[0]
        cod, nome = self.tree_cad.item(s[0])["values"]
        self.v_cod.set(str(cod))
        self.v_nome.set(str(nome))
        self._mat_nome = str(nome)
        self.lbl_mat.config(text=f"Editando: {nome}", fg=AZUL)

    def mat_limpar(self):
        self.mat_sel = None
        self._mat_nome = None
        self.v_cod.set("")
        self.v_nome.set("")
        self.tree_cad.selection_remove(*self.tree_cad.selection())
        self.lbl_mat.config(text="Preencha o código e o nome e clique em \"Adicionar novo\".", fg=CINZA)

    def _campos(self):
        nome = re.sub(r"\s+", " ", self.v_nome.get()).strip()
        cod = self.v_cod.get().strip()
        if not nome:
            messagebox.showinfo("Cadastro", "Informe o nome do material.")
            return None
        if cod and not re.fullmatch(r"[0-9A-Za-z.\-/]+", cod):
            messagebox.showinfo("Cadastro", "Código inválido.")
            return None
        return cod, nome

    def _feito(self, msg, nome=None):
        self.app.cadastro_mudou()
        if nome is not None:
            k = chave_nome(nome)
            i = next((i for i, m in enumerate(CADASTRO.materiais) if chave_nome(m) == k), None)
            self.mat_sel = f"m{i}" if i is not None else None
            self._mat_nome = nome
        self.mostrar_cadastro()
        self.app.msg(msg)
        self.lbl_mat.config(text=msg, fg=VERDE)

    def mat_adicionar(self):
        c = self._campos()
        if not c:
            return
        try:
            CADASTRO.adicionar_material(*c)
        except ValueError as ex:
            messagebox.showinfo("Cadastro", str(ex))
            return
        self.v_busca.set("")
        self._feito(f"Adicionado: {c[0]} {c[1]}", c[1])

    def mat_salvar(self):
        if not getattr(self, "_mat_nome", None):
            messagebox.showinfo("Cadastro", "Clique na lista no material que quer alterar "
                                            "(ou use \"Adicionar novo\").")
            return
        c = self._campos()
        if not c:
            return
        try:
            CADASTRO.editar_material(self._mat_nome, *c)
        except ValueError as ex:
            messagebox.showinfo("Cadastro", str(ex))
            return
        self._feito(f"Alterado: {c[0]} {c[1]}", c[1])

    def mat_excluir(self):
        nome = getattr(self, "_mat_nome", None)
        if not nome or not self.app.pode_editar:
            if not nome:
                messagebox.showinfo("Cadastro", "Clique na lista no material que quer excluir.")
            return
        if not messagebox.askyesno("Cadastro", f"Excluir \"{nome}\" do cadastro?"):
            return
        CADASTRO.excluir_material(nome)
        self.mat_limpar()
        self._feito(f"Excluído: {nome}")

    def baixar_cadastro(self):
        cam = filedialog.asksaveasfilename(defaultextension=".xlsx", filetypes=[("Excel", "*.xlsx")],
                                           initialfile="cadastro_materiais.xlsx")
        if not cam:
            return
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill
            wb = Workbook()
            ws = wb.active
            ws.title = "Materiais"
            ws.append(["Código", "Material"])
            for c in ws[1]:
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = PatternFill("solid", fgColor="014FA0")
            for m in CADASTRO.materiais:
                cod = CADASTRO.codigo_de(m)
                ws.append([int(cod) if cod.isdigit() else cod, m])
            ws.column_dimensions["A"].width = 14
            ws.column_dimensions["B"].width = 60
            wb.save(cam)
        except Exception as ex:
            messagebox.showerror("Baixar cadastro", f"{ex}\n\nSe a planilha estiver aberta no Excel, feche e "
                                                    "tente de novo.")
            return
        messagebox.showinfo("Baixar cadastro", f"Cadastro salvo:\n{cam}")

    def _nome_sel(self):
        s = [i for i in self.tree_map.selection() if i.startswith("n:")]
        return s[0][2:] if s else None

    def escolher(self):
        if not self.app.pode_editar:
            return
        nome = self._nome_sel()
        if nome is None:
            messagebox.showinfo("Cadastro", "Selecione um nome na lista.")
            return
        if not CADASTRO.materiais:
            messagebox.showinfo("Cadastro", "Carregue a planilha de materiais primeiro.")
            return
        r = EscolherNome(self, nome).resultado
        if r is not None:
            CADASTRO.ajustar(nome, r)
            self.app.cadastro_mudou()

    def automatico(self):
        if not self.app.pode_editar:
            return
        nome = self._nome_sel()
        if nome is not None:
            CADASTRO.ajustar(nome, None)
            self.app.cadastro_mudou()

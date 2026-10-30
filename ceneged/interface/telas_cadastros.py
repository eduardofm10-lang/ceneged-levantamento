"""Telas de equipes, obras, constantes e carteira.

Parte do programa Levantamento de Materiais (Ceneged)."""

import time
import tkinter as tk
from tkinter import filedialog
from tkinter import messagebox
from tkinter import ttk

from ceneged.comum.texto import (chave_nome, formata_num, normaliza_codigo)
from ceneged.dados.equipes import (CAMPOS_OBRA, CARTEIRA, CORES_STATUS, EQUIPES, FUNCOES,
    MESES, OPC_SN, OPC_SNP, STATUS_OBRA, data_br, data_ordem)
from ceneged.dados.constantes import (CONSTANTES, _num_br)
from ceneged.dados.planilhas import (ler_planilha_carteira, ler_planilha_constantes,
    salvar_planilha_carteira)
from ceneged.dominio.materiais import (achar_material_parecido, ligar_constantes,
    materiais_conhecidos)
from ceneged.interface.tema import (AZUL, BORDA, BRANCO, CINZA, FONTE, FUNDO, LARANJA, TEXTO,
    VERDE, VERMELHO, cartao, secao, tabela)
from ceneged.interface.dialogos import EscolherMaterial

class _TelaBase(tk.Frame):
    def _ed(self, w):
        self.__dict__.setdefault("edicao", []).append(w)
        return w

    def aplicar_permissao(self, pode):
        self.pode = pode
        for w in getattr(self, "edicao", []):
            try:
                w.state(["!disabled"] if pode else ["disabled"])
            except tk.TclError:
                pass
        if not pode and not hasattr(self, "_aviso_leitura"):
            self._aviso_leitura = tk.Label(self, text="Você pode consultar, mas não alterar. "
                                                      "Fale com o administrador.",
                                           bg="#FFF4D6", fg="#7A5600", font=(FONTE, 9), anchor="w", padx=10, pady=6)
            self._aviso_leitura.pack(fill="x", padx=20, pady=(0, 8), after=self.winfo_children()[0])

    def _cabeca(self, titulo, sub):
        cab = tk.Frame(self, bg=FUNDO)
        cab.pack(fill="x", padx=20, pady=(14, 8))
        ttk.Label(cab, text=titulo, style="Titulo.TLabel").pack(anchor="w")
        ttk.Label(cab, style="Sub.TLabel", wraplength=1150, justify="left", text=sub).pack(anchor="w")


def _rotulo(pai, texto, row, col, **kw):
    tk.Label(pai, text=texto, bg=BRANCO, fg=CINZA, font=(FONTE, 8)).grid(row=row, column=col, sticky="w", **kw)


class TelaEquipes(_TelaBase):
    """Pessoas (supervisores, encarregados, colaboradores) e equipes."""
    FILTROS = ("Todas as funções",) + FUNCOES

    def __init__(self, pai, app):
        super().__init__(pai, bg=FUNDO)
        self.app = app
        self.pode = True
        self.pessoa_sel = None
        self.equipe_sel = None
        self._cabeca("Cadastro de equipe",
                     "Cadastre as pessoas (supervisores, encarregados e colaboradores) e depois monte cada equipe "
                     "com o código, o supervisor, o encarregado e os colaboradores. Eles aparecem como opção na "
                     "Carteira de obra.")
        corpo = tk.Frame(self, bg=FUNDO)
        corpo.pack(fill="both", expand=True, padx=20, pady=(0, 16))
        corpo.columnconfigure(0, weight=2, uniform="e")
        corpo.columnconfigure(1, weight=3, uniform="e")
        corpo.rowconfigure(0, weight=1)

        # ---------------- pessoas
        esq = cartao(corpo)
        esq.grid(row=0, column=0, sticky="nsew", padx=(0, 10))
        self.sec_p = secao(esq, "Pessoas")
        self.sec_p.pack(anchor="w", padx=14, pady=(12, 6))
        fb = tk.Frame(esq, bg=BRANCO)
        fb.pack(fill="x", padx=14)
        self.v_filtro = tk.StringVar(value=self.FILTROS[0])
        cb = ttk.Combobox(fb, textvariable=self.v_filtro, values=self.FILTROS, state="readonly", width=16)
        cb.pack(side="left")
        cb.bind("<<ComboboxSelected>>", lambda e: self.mostrar_pessoas())
        self.v_busca_p = tk.StringVar()
        ttk.Entry(fb, textvariable=self.v_busca_p).pack(side="left", fill="x", expand=True, padx=(6, 0))
        self.v_busca_p.trace_add("write", lambda *a: self.mostrar_pessoas())
        # formulário (embaixo, sempre visível)
        fp = tk.Frame(esq, bg=BRANCO)
        fp.pack(side="bottom", fill="x", padx=14, pady=(6, 12))
        self.vp_nome, self.vp_funcao, self.vp_mat = tk.StringVar(), tk.StringVar(value=FUNCOES[2]), tk.StringVar()
        _rotulo(fp, "Nome", 0, 0)
        _rotulo(fp, "Função", 0, 1, padx=(6, 0))
        _rotulo(fp, "Matrícula (opcional)", 0, 2, padx=(6, 0))
        self._ed(ttk.Entry(fp, textvariable=self.vp_nome)).grid(row=1, column=0, sticky="we")
        self._ed(ttk.Combobox(fp, textvariable=self.vp_funcao, values=FUNCOES, state="readonly",
                              width=12)).grid(row=1, column=1, padx=(6, 0))
        self._ed(ttk.Entry(fp, textvariable=self.vp_mat, width=12)).grid(row=1, column=2, padx=(6, 0))
        fp.columnconfigure(0, weight=1)
        bp = tk.Frame(fp, bg=BRANCO)
        bp.grid(row=2, column=0, columnspan=3, sticky="we", pady=(6, 0))
        self._ed(ttk.Button(bp, text="Adicionar pessoa", style="Azul.TButton",
                            command=self.p_adicionar)).pack(side="left")
        self._ed(ttk.Button(bp, text="Salvar alteração", command=self.p_salvar)).pack(side="left", padx=4)
        self._ed(ttk.Button(bp, text="Excluir", command=self.p_excluir)).pack(side="left")
        ttk.Button(bp, text="Limpar", command=self.p_limpar).pack(side="right")
        self.lbl_p = tk.Label(fp, text="", bg=BRANCO, fg=CINZA, font=(FONTE, 8), anchor="w")
        self.lbl_p.grid(row=3, column=0, columnspan=3, sticky="we", pady=(4, 0))
        self.tree_p = tabela(esq, [("nome", "Nome", 170), ("funcao", "Função", 95), ("eq", "Equipe", 80)], altura=4)
        self.tree_p.quadro.pack(fill="both", expand=True, padx=14, pady=(6, 0))
        self.tree_p.bind("<<TreeviewSelect>>", lambda e: self.p_abrir())
        for f, cor in zip(FUNCOES, (AZUL, VERDE, TEXTO)):
            self.tree_p.tag_configure(f, foreground=cor)

        # ---------------- equipes
        dir_ = cartao(corpo)
        dir_.grid(row=0, column=1, sticky="nsew")
        self.sec_e = secao(dir_, "Equipes")
        self.sec_e.pack(anchor="w", padx=14, pady=(12, 6))
        fe = tk.Frame(dir_, bg=BRANCO)
        fe.pack(side="bottom", fill="x", padx=14, pady=(6, 12))
        self.ve_cod, self.ve_sup, self.ve_enc = tk.StringVar(), tk.StringVar(), tk.StringVar()
        _rotulo(fe, "Código da equipe", 0, 0)
        _rotulo(fe, "Supervisor", 0, 1, padx=(6, 0))
        _rotulo(fe, "Encarregado", 0, 2, padx=(6, 0))
        self._ed(ttk.Entry(fe, textvariable=self.ve_cod, width=12)).grid(row=1, column=0, sticky="we")
        self.cb_sup = self._ed(ttk.Combobox(fe, textvariable=self.ve_sup, state="readonly"))
        self.cb_sup.grid(row=1, column=1, sticky="we", padx=(6, 0))
        self.cb_enc = self._ed(ttk.Combobox(fe, textvariable=self.ve_enc, state="readonly"))
        self.cb_enc.grid(row=1, column=2, sticky="we", padx=(6, 0))
        fe.columnconfigure(1, weight=1)
        fe.columnconfigure(2, weight=1)
        col = tk.Frame(fe, bg=BRANCO)
        col.grid(row=2, column=0, columnspan=3, sticky="we", pady=(8, 0))
        col.columnconfigure(0, weight=1)
        col.columnconfigure(2, weight=1)
        _rotulo(col, "Colaboradores disponíveis", 0, 0)
        _rotulo(col, "Colaboradores desta equipe", 0, 2)
        lst_kw = dict(height=5, bg=BRANCO, fg=TEXTO, selectbackground=AZUL, selectforeground=BRANCO,
                      relief="flat", highlightthickness=1, highlightbackground=BORDA, font=(FONTE, 9),
                      exportselection=False, selectmode="extended")
        self.lst_disp = tk.Listbox(col, **lst_kw)
        self.lst_disp.grid(row=1, column=0, sticky="nsew")
        self.lst_eq = tk.Listbox(col, **lst_kw)
        self.lst_eq.grid(row=1, column=2, sticky="nsew")
        mv = tk.Frame(col, bg=BRANCO)
        mv.grid(row=1, column=1, padx=8)
        self._ed(ttk.Button(mv, text="Colocar >>", command=self.e_colocar)).pack(pady=(0, 4))
        self._ed(ttk.Button(mv, text="<< Tirar", command=self.e_tirar)).pack()
        self.lst_disp.bind("<Double-1>", lambda e: self.e_colocar())
        self.lst_eq.bind("<Double-1>", lambda e: self.e_tirar())
        be = tk.Frame(fe, bg=BRANCO)
        be.grid(row=3, column=0, columnspan=3, sticky="we", pady=(8, 0))
        self._ed(ttk.Button(be, text="Nova equipe", command=self.e_nova)).pack(side="left")
        self._ed(ttk.Button(be, text="Excluir equipe", command=self.e_excluir)).pack(side="left", padx=4)
        self._ed(ttk.Button(be, text="Salvar equipe", style="Verm.TButton",
                            command=self.e_salvar)).pack(side="right")
        self.lbl_e = tk.Label(fe, text="", bg=BRANCO, fg=CINZA, font=(FONTE, 8), anchor="w")
        self.lbl_e.grid(row=4, column=0, columnspan=3, sticky="we", pady=(4, 0))
        self.tree_e = tabela(dir_, [("cod", "Código", 80), ("sup", "Supervisor", 120), ("enc", "Encarregado", 120),
                                    ("col", "Colaboradores", 260)], altura=4)
        self.tree_e.quadro.pack(fill="both", expand=True, padx=14)
        self.tree_e.bind("<<TreeviewSelect>>", lambda e: self.e_abrir())
        self.colab = []
        self.e_nova(limpar_sel=False)

    # ---------------------------------------------------------------- listas
    def atualizar(self):
        EQUIPES.recarregar_se_mudou()
        self.cb_sup.configure(values=[""] + EQUIPES.nomes("Supervisor"))
        self.cb_enc.configure(values=[""] + EQUIPES.nomes("Encarregado"))
        self.mostrar_pessoas()
        self.mostrar_equipes()
        self.mostrar_colab()

    def mostrar_pessoas(self):
        t = self.tree_p
        t.delete(*t.get_children())
        f = self.v_filtro.get()
        k = chave_nome(self.v_busca_p.get())
        for i, p in enumerate(sorted(EQUIPES.pessoas, key=lambda p: (FUNCOES.index(p["funcao"])
                                                                   if p["funcao"] in FUNCOES else 9,
                                                                   p["nome"]))):
            if f in FUNCOES and p["funcao"] != f:
                continue
            if k and k not in chave_nome(p["nome"] + p.get("matricula", "")):
                continue
            t.insert("", "end", iid="p:" + p["nome"], values=(p["nome"], p["funcao"],
                                                              ", ".join(EQUIPES.equipes_da_pessoa(p["nome"]))),
                     tags=(p["funcao"],))
        if self.pessoa_sel and t.exists("p:" + self.pessoa_sel):
            t.selection_set("p:" + self.pessoa_sel)
            t.see("p:" + self.pessoa_sel)
        n = {f_: sum(1 for p in EQUIPES.pessoas if p["funcao"] == f_) for f_ in FUNCOES}
        self.sec_p.rotulo.config(text=f"Pessoas ({len(EQUIPES.pessoas)})  ·  {n['Supervisor']} supervisor(es), "
                                      f"{n['Encarregado']} encarregado(s), {n['Colaborador']} colaborador(es)")

    def mostrar_equipes(self):
        t = self.tree_e
        t.delete(*t.get_children())
        for c in EQUIPES.codigos():
            e = EQUIPES.equipe(c)
            t.insert("", "end", iid="e:" + c, values=(c, e.get("supervisor", ""), e.get("encarregado", ""),
                                                      ", ".join(e.get("colaboradores", []))))
        if self.equipe_sel and t.exists("e:" + self.equipe_sel):
            t.selection_set("e:" + self.equipe_sel)
            t.see("e:" + self.equipe_sel)
        self.sec_e.rotulo.config(text=f"Equipes ({len(EQUIPES.equipes)})")

    def mostrar_colab(self):
        dentro = {chave_nome(c) for c in self.colab}
        self.lst_disp.delete(0, "end")
        self._disp = []
        for n in EQUIPES.nomes("Colaborador"):
            if chave_nome(n) in dentro:
                continue
            outras = [c for c in EQUIPES.equipes_da_pessoa(n) if c != self.equipe_sel]
            self._disp.append(n)
            self.lst_disp.insert("end", n + (f"   ({', '.join(outras)})" if outras else ""))
            if outras:
                self.lst_disp.itemconfig("end", fg=CINZA)
        self.lst_eq.delete(0, "end")
        for n in self.colab:
            self.lst_eq.insert("end", n)

    # ---------------------------------------------------------------- pessoas
    def _msg(self, lbl, texto, cor=VERDE):
        lbl.config(text=texto, fg=cor)
        self.app.msg(texto)

    def p_abrir(self):
        s = self.tree_p.selection()
        if not s:
            return
        p = EQUIPES.pessoa(s[0][2:])
        if not p:
            return
        self.pessoa_sel = p["nome"]
        self.vp_nome.set(p["nome"])
        self.vp_funcao.set(p["funcao"])
        self.vp_mat.set(p.get("matricula", ""))
        self.lbl_p.config(text=f"Editando: {p['nome']}", fg=AZUL)

    def p_limpar(self):
        self.pessoa_sel = None
        self.vp_nome.set("")
        self.vp_mat.set("")
        self.tree_p.selection_remove(*self.tree_p.selection())
        self.lbl_p.config(text="Preencha o nome e a função e clique em \"Adicionar pessoa\".", fg=CINZA)

    def _p_gravar(self, antigo):
        try:
            nome = EQUIPES.salvar_pessoa(antigo, self.vp_nome.get(), self.vp_funcao.get(), self.vp_mat.get())
        except ValueError as ex:
            messagebox.showinfo("Pessoas", str(ex))
            return
        self.pessoa_sel = nome
        self.atualizar()
        self._msg(self.lbl_p, f"{'Alterado' if antigo else 'Cadastrado'}: {nome} ({self.vp_funcao.get()})")

    def p_adicionar(self):
        self._p_gravar(None)

    def p_salvar(self):
        if not self.pessoa_sel:
            messagebox.showinfo("Pessoas", "Clique na lista na pessoa que quer alterar (ou use \"Adicionar pessoa\").")
            return
        self._p_gravar(self.pessoa_sel)

    def p_excluir(self):
        if not self.pessoa_sel:
            messagebox.showinfo("Pessoas", "Clique na lista na pessoa que quer excluir.")
            return
        eqs = EQUIPES.equipes_da_pessoa(self.pessoa_sel)
        if not messagebox.askyesno("Pessoas", f"Excluir {self.pessoa_sel}?" +
                                   (f"\n\nEla sai também da(s) equipe(s): {', '.join(eqs)}." if eqs else "")):
            return
        nome = self.pessoa_sel
        EQUIPES.excluir_pessoa(nome)
        self.p_limpar()
        self.atualizar()
        self._msg(self.lbl_p, f"Excluído: {nome}")

    # ---------------------------------------------------------------- equipes
    def e_abrir(self):
        s = self.tree_e.selection()
        if not s:
            return
        e = EQUIPES.equipe(s[0][2:])
        if not e:
            return
        self.equipe_sel = e["codigo"]
        self.ve_cod.set(e["codigo"])
        self.ve_sup.set(e.get("supervisor", ""))
        self.ve_enc.set(e.get("encarregado", ""))
        self.colab = list(e.get("colaboradores", []))
        self.mostrar_colab()
        self.lbl_e.config(text=f"Editando a equipe {e['codigo']}", fg=AZUL)

    def e_nova(self, limpar_sel=True):
        self.equipe_sel = None
        self.ve_cod.set("")
        self.ve_sup.set("")
        self.ve_enc.set("")
        self.colab = []
        if limpar_sel:
            self.tree_e.selection_remove(*self.tree_e.selection())
        self.mostrar_colab()
        self.lbl_e.config(text="Nova equipe: preencha o código, escolha o supervisor, o encarregado e os "
                               "colaboradores e clique em \"Salvar equipe\".", fg=CINZA)

    def e_colocar(self):
        for i in self.lst_disp.curselection():
            self.colab.append(self._disp[i])
        self.mostrar_colab()

    def e_tirar(self):
        tirar = {self.lst_eq.get(i) for i in self.lst_eq.curselection()}
        self.colab = [c for c in self.colab if c not in tirar]
        self.mostrar_colab()

    def e_salvar(self):
        try:
            cod = EQUIPES.salvar_equipe(self.equipe_sel, self.ve_cod.get(), self.ve_sup.get(), self.ve_enc.get(),
                                        self.colab)
        except ValueError as ex:
            messagebox.showinfo("Equipes", str(ex))
            return
        novo = self.equipe_sel is None
        self.equipe_sel = cod
        self.atualizar()
        self._msg(self.lbl_e, f"Equipe {cod} {'cadastrada' if novo else 'salva'}.")

    def e_excluir(self):
        if not self.equipe_sel:
            messagebox.showinfo("Equipes", "Clique na lista na equipe que quer excluir.")
            return
        if not messagebox.askyesno("Equipes", f"Excluir a equipe {self.equipe_sel}?\n\n"
                                              "As pessoas continuam cadastradas."):
            return
        cod = self.equipe_sel
        EQUIPES.excluir_equipe(cod)
        self.e_nova()
        self.atualizar()
        self._msg(self.lbl_e, f"Equipe {cod} excluída.")


class FormObra(tk.Toplevel):
    """Formulário de um lançamento da carteira de obra (novo ou edição)."""

    def __init__(self, pai, obra=None, usuario=""):
        super().__init__(pai)
        self.obra = obra
        self.usuario = usuario
        self.resultado = None
        titulo = "Novo lançamento" if obra is None else f"Editar lançamento — projeto {obra.get('projeto', '')}"
        self.title(titulo)
        self.configure(bg=BRANCO)
        self.transient(pai.winfo_toplevel())
        self.resizable(False, False)
        tk.Frame(self, bg=VERMELHO, height=3).pack(fill="x")
        tk.Label(self, text=titulo, bg=BRANCO, fg=AZUL, font=(FONTE, 14, "bold")).pack(anchor="w", padx=18,
                                                                                       pady=(12, 0))
        tk.Label(self, text="Datas no formato dd/mm/aaaa (ex.: 05/10/2026). Campos com * são obrigatórios.",
                 bg=BRANCO, fg=CINZA, font=(FONTE, 9)).pack(anchor="w", padx=18, pady=(0, 8))
        f = tk.Frame(self, bg=BRANCO)
        f.pack(fill="both", padx=18)
        o = obra or {}
        self.v = {k: tk.StringVar(value=o.get(k, "")) for k, _, _ in CAMPOS_OBRA if k != "obs"}
        if obra is None:
            self.v["viabilizada"].set("SIM")
            self.v["requisicao"].set("SIM")
            self.v["desligamento"].set("NÃO")
            self.v["status"].set("EM ANDAMENTO")

        def campo(row, col, rotulo, w):
            tk.Label(f, text=rotulo, bg=BRANCO, fg=TEXTO, font=(FONTE, 9, "bold")).grid(
                row=row * 2, column=col, sticky="w", padx=(0, 14), pady=(6, 0))
            w.grid(row=row * 2 + 1, column=col, sticky="we", padx=(0, 14))
            return w

        def combo(k, valores, ler=True, largura=22):
            return ttk.Combobox(f, textvariable=self.v[k], values=valores, state="readonly" if ler else "normal",
                                width=largura)

        def data(k):
            fr = tk.Frame(f, bg=BRANCO)
            e = ttk.Entry(fr, textvariable=self.v[k], width=14)
            e.pack(side="left")
            ttk.Button(fr, text="Hoje", style="Peq.TButton",
                       command=lambda: self.v[k].set(time.strftime("%d/%m/%Y"))).pack(side="left", padx=(4, 0))
            fr.entrada = e
            return fr

        def pessoa(k, funcao):
            fr = tk.Frame(f, bg=BRANCO)
            cb = ttk.Combobox(fr, textvariable=self.v[k], values=[""] + EQUIPES.nomes(funcao), state="readonly",
                              width=20)
            cb.pack(side="left", fill="x", expand=True)
            ttk.Button(fr, text="+", width=3, style="Peq.TButton",
                       command=lambda: self._nova_pessoa(k, funcao, cb)).pack(side="left", padx=(4, 0))
            return fr

        self.e_proj = campo(0, 0, "Projeto *", ttk.Entry(f, textvariable=self.v["projeto"], width=22))
        campo(0, 1, "Nota", ttk.Entry(f, textvariable=self.v["nota"], width=22))
        cb_eq = campo(0, 2, "Equipe", combo("equipe", [""] + EQUIPES.codigos()))
        cb_eq.bind("<<ComboboxSelected>>", lambda e: self._equipe_escolhida())
        campo(1, 0, "Supervisor", pessoa("supervisor", "Supervisor"))
        campo(1, 1, "Encarregado", pessoa("encarregado", "Encarregado"))
        campo(1, 2, "Status da obra", combo("status", STATUS_OBRA))
        campo(2, 0, "Viabilizada?", combo("viabilizada", OPC_SNP))
        campo(2, 1, "Requisição", combo("requisicao", OPC_SNP))
        campo(3, 0, "Data de início", data("inicio"))
        campo(3, 1, "Data de finalização", data("fim"))
        campo(3, 2, "Data de conclusão", data("conclusao"))
        cb_d = campo(4, 0, "Desligamento", combo("desligamento", OPC_SN))
        cb_d.bind("<<ComboboxSelected>>", lambda e: self._deslig())
        self.e_desl = campo(4, 1, "Data do desligamento (pode ser período: 05/10/2026 - 06/10/2026)",
                            ttk.Entry(f, textvariable=self.v["data_desl"], width=40))
        self.e_desl.grid(columnspan=2)
        tk.Label(f, text="Observações", bg=BRANCO, fg=TEXTO, font=(FONTE, 9, "bold")).grid(
            row=10, column=0, sticky="w", pady=(6, 0))
        self.t_obs = tk.Text(f, height=4, width=80, font=(FONTE, 10), relief="flat", highlightthickness=1,
                             highlightbackground=BORDA, highlightcolor=AZUL, wrap="word")
        self.t_obs.grid(row=11, column=0, columnspan=3, sticky="we", padx=(0, 14))
        self.t_obs.insert("1.0", o.get("obs", ""))
        if obra:
            info = f"Lançado por {obra.get('criado_por') or '—'} em {obra.get('criado_em') or '—'}"
            if obra.get("alterado_em"):
                info += f"  ·  última alteração: {obra.get('alterado_por') or '—'} em {obra['alterado_em']}"
            tk.Label(self, text=info, bg=BRANCO, fg=CINZA, font=(FONTE, 8)).pack(anchor="w", padx=18, pady=(8, 0))
        bts = tk.Frame(self, bg=BRANCO)
        bts.pack(fill="x", padx=18, pady=14)
        ttk.Button(bts, text="Concluído", style="Verm.TButton", command=self.ok).pack(side="right")
        ttk.Button(bts, text="Cancelar", command=self.destroy).pack(side="right", padx=6)
        self._deslig()
        self.bind("<Escape>", lambda e: self.destroy())
        self.geometry("+%d+%d" % (pai.winfo_rootx() + 120, pai.winfo_rooty() + 30))
        self.e_proj.focus_set()
        self.grab_set()
        self.wait_window()

    def _nova_pessoa(self, k, funcao, cb):
        from tkinter import simpledialog
        n = simpledialog.askstring(f"Novo {funcao.lower()}", f"Nome do {funcao.lower()}:", parent=self)
        if not n or not n.strip():
            return
        try:
            nome = EQUIPES.salvar_pessoa(None, n, funcao)
        except ValueError as ex:
            messagebox.showinfo("Cadastro", str(ex), parent=self)
            return
        cb.configure(values=[""] + EQUIPES.nomes(funcao))
        self.v[k].set(nome)

    def _equipe_escolhida(self):
        e = EQUIPES.equipe(self.v["equipe"].get())
        if e:
            if e.get("supervisor"):
                self.v["supervisor"].set(e["supervisor"])
            if e.get("encarregado"):
                self.v["encarregado"].set(e["encarregado"])

    def _deslig(self):
        if self.v["desligamento"].get() == "SIM":
            self.e_desl.state(["!disabled"])
        else:
            self.v["data_desl"].set("")
            self.e_desl.state(["disabled"])

    def ok(self):
        d = {k: v.get().strip() for k, v in self.v.items()}
        d["obs"] = self.t_obs.get("1.0", "end").strip()
        if not d["projeto"]:
            messagebox.showinfo("Lançamento", "Informe o projeto.", parent=self)
            return
        try:
            for k in ("inicio", "fim", "conclusao"):
                d[k] = data_br(d[k])
        except ValueError as ex:
            messagebox.showinfo("Lançamento", str(ex), parent=self)
            return
        if d["inicio"] and d["fim"] and data_ordem(d["fim"]) < data_ordem(d["inicio"]):
            messagebox.showinfo("Lançamento", "A data de finalização está antes da data de início.", parent=self)
            return
        if d["desligamento"] == "SIM" and not d["data_desl"]:
            if not messagebox.askyesno("Lançamento", "Desligamento = SIM, mas sem a data do desligamento.\n\n"
                                                     "Salvar mesmo assim?", parent=self):
                return
        if d["status"] == "CONCLUÍDA" and not d["conclusao"]:
            if not messagebox.askyesno("Lançamento", "Status CONCLUÍDA, mas sem a data de conclusão.\n\n"
                                                     "Salvar mesmo assim?", parent=self):
                return
        dup = CARTEIRA.duplicada(d["projeto"], d["nota"], self.obra.get("id") if self.obra else None)
        if dup and not messagebox.askyesno("Lançamento", f"Já existe um lançamento do projeto {d['projeto']}"
                                                         f"{' / nota ' + d['nota'] if d['nota'] else ''}.\n\n"
                                                         "Salvar outro mesmo assim?", parent=self):
            return
        try:
            self.resultado = CARTEIRA.salvar_obra(d, self.obra.get("id") if self.obra else None, self.usuario)
        except ValueError as ex:
            messagebox.showinfo("Lançamento", str(ex), parent=self)
            return
        self.destroy()


class TelaConstantes(_TelaBase):
    """Constantes de metro para quilo (cabos nus etc.)."""

    def __init__(self, pai, app):
        super().__init__(pai, bg=FUNDO)
        self.app = app
        self.pode = True
        self.sel = None
        self._cabeca("Constantes (metro → KG)",
                     "Materiais lançados em metros que vão para o relatório em quilos: quantidade em metros × "
                     "constante = KG. O material é reconhecido pelo código ou pelo nome igual ao do relatório.")
        bar = cartao(self)
        bar.pack(fill="x", padx=20, pady=(0, 8))
        b = tk.Frame(bar, bg=BRANCO)
        b.pack(fill="x", padx=12, pady=10)
        self._ed(ttk.Button(b, text="Carregar planilha", style="Azul.TButton", command=self.carregar)).pack(side="left")
        ttk.Button(b, text="Baixar planilha", command=self.baixar).pack(side="left", padx=6)
        self._ed(ttk.Button(b, text="Usar os nomes do cadastro", command=self.ligar)).pack(side="left")
        self.lbl_n = tk.Label(b, text="", bg=BRANCO, fg=CINZA, font=(FONTE, 9))
        self.lbl_n.pack(side="left", padx=10)

        c = cartao(self)
        c.pack(fill="both", expand=True, padx=20, pady=(0, 16))
        f = tk.Frame(c, bg=BRANCO)
        f.pack(side="bottom", fill="x", padx=14, pady=(6, 12))
        self.v_cod, self.v_desc, self.v_k = tk.StringVar(), tk.StringVar(), tk.StringVar()
        _rotulo(f, "Código (opcional)", 0, 0)
        _rotulo(f, "Descrição do material (igual ao nome que sai no relatório)", 0, 1, padx=(6, 0))
        _rotulo(f, "Constante (kg por metro)", 0, 3, padx=(6, 0))
        self._ed(ttk.Entry(f, textvariable=self.v_cod, width=12)).grid(row=1, column=0)
        self._ed(ttk.Entry(f, textvariable=self.v_desc)).grid(row=1, column=1, sticky="we", padx=(6, 0))
        self._ed(ttk.Button(f, text="Procurar material…", command=self.procurar)).grid(row=1, column=2, padx=(6, 0))
        self._ed(ttk.Entry(f, textvariable=self.v_k, width=12)).grid(row=1, column=3, padx=(6, 0))
        f.columnconfigure(1, weight=1)
        bt = tk.Frame(f, bg=BRANCO)
        bt.grid(row=2, column=0, columnspan=4, sticky="we", pady=(8, 0))
        self._ed(ttk.Button(bt, text="Adicionar", style="Azul.TButton", command=self.adicionar)).pack(side="left")
        self._ed(ttk.Button(bt, text="Salvar alteração", command=self.salvar)).pack(side="left", padx=4)
        self._ed(ttk.Button(bt, text="Excluir", command=self.excluir)).pack(side="left")
        ttk.Button(bt, text="Limpar campos", command=self.limpar).pack(side="right")
        self.lbl = tk.Label(f, text="", bg=BRANCO, fg=CINZA, font=(FONTE, 8), anchor="w")
        self.lbl.grid(row=3, column=0, columnspan=4, sticky="we", pady=(4, 0))
        secao(c, "Materiais com constante").pack(anchor="w", padx=14, pady=(12, 6))
        self.tree = tabela(c, [("cod", "Código", 90), ("desc", "Descrição", 360), ("k", "Constante (kg/m)", 120),
                               ("ex", "Exemplo: 100 m =", 120), ("ach", "No cadastro / materiais extras", 260)],
                           altura=6)
        self.tree.quadro.pack(fill="both", expand=True, padx=14)
        self.tree.tag_configure("nao", foreground=LARANJA)
        self.tree.bind("<<TreeviewSelect>>", lambda e: self.abrir())
        self.tree.bind("<Delete>", lambda e: self.excluir())
        self.limpar()

    def atualizar(self):
        CONSTANTES.recarregar_se_mudou()
        conhecidos = materiais_conhecidos()
        por_nome = {chave_nome(n): (c, n) for c, n, _ in conhecidos.values()}
        por_cod = {c: (c, n) for c, n, _ in conhecidos.values() if c}
        t = self.tree
        t.delete(*t.get_children())
        for j, i in enumerate(CONSTANTES.itens):
            ach = por_cod.get(normaliza_codigo(i.get("codigo", ""))) or por_nome.get(chave_nome(i["descricao"]))
            txt = f"{ach[0] or 's/ código'} — {ach[1]}" if ach else "não encontrado (use Procurar material…)"
            if i.get("planilha"):
                txt += f"   (na planilha: {i['planilha']})"
            t.insert("", "end", iid=str(j), values=(i.get("codigo", ""), i["descricao"],
                                                    _num_br(i["constante"]),
                                                    formata_num(100 * i["constante"]) + " KG", txt),
                     tags=() if ach else ("nao",))
        if self.sel is not None and t.exists(str(self.sel)):
            t.selection_set(str(self.sel))
        self.lbl_n.config(text=f"{len(CONSTANTES.itens)} material(is) com constante")

    def ligar(self):
        CONSTANTES.recarregar_se_mudou()
        antes = [i["descricao"] for i in CONSTANTES.itens]
        ligar_constantes(CONSTANTES.itens)
        CONSTANTES.salvar()
        trocados = sum(1 for a, i in zip(antes, CONSTANTES.itens) if a != i["descricao"])
        sem = sum(1 for i in CONSTANTES.itens if not achar_material_parecido(i["descricao"]))
        self.app.cadastro_mudou()
        self.atualizar()
        messagebox.showinfo("Constantes", f"{trocados} nome(s) trocado(s) pelo nome do material do sistema." +
                            (f"\n\n{sem} não encontrado(s) (em laranja): clique na linha, use Procurar "
                             "material… e Salvar alteração." if sem else ""))

    def abrir(self):
        s = self.tree.selection()
        if not s:
            return
        self.sel = int(s[0])
        i = CONSTANTES.itens[self.sel]
        self.v_cod.set(i.get("codigo", ""))
        self.v_desc.set(i["descricao"])
        self.v_k.set(_num_br(i["constante"]))
        self.lbl.config(text=f"Editando: {i['descricao']}", fg=AZUL)

    def limpar(self):
        self.sel = None
        for v in (self.v_cod, self.v_desc, self.v_k):
            v.set("")
        self.tree.selection_remove(*self.tree.selection())
        self.lbl.config(text="Preencha a descrição (ou use Procurar material…) e a constante e clique em Adicionar.",
                        fg=CINZA)

    def procurar(self):
        r = EscolherMaterial(self, "Escolher material", com_kits=False, com_qtd=False).resultado
        if r:
            self.v_cod.set(r["codigo"])
            self.v_desc.set(r["nome"])

    def _gravar(self, indice):
        try:
            CONSTANTES.salvar_item(indice, self.v_cod.get(), self.v_desc.get(), self.v_k.get())
        except ValueError as ex:
            messagebox.showinfo("Constantes", str(ex))
            return
        self.sel = indice if indice is not None else len(CONSTANTES.itens) - 1
        self.app.cadastro_mudou()
        self.atualizar()
        self.lbl.config(text=f"{'Alterado' if indice is not None else 'Adicionado'}: {self.v_desc.get().strip()}",
                        fg=VERDE)

    def adicionar(self):
        self._gravar(None)

    def salvar(self):
        if self.sel is None:
            messagebox.showinfo("Constantes", "Clique na lista no material que quer alterar (ou use Adicionar).")
            return
        self._gravar(self.sel)

    def excluir(self):
        if self.sel is None or not self.pode:
            return
        i = CONSTANTES.itens[self.sel]
        if not messagebox.askyesno("Constantes", f"Excluir a constante de {i['descricao']}?\n\n"
                                                 "Ele volta a sair em metros no relatório."):
            return
        CONSTANTES.excluir_item(self.sel)
        self.limpar()
        self.app.cadastro_mudou()
        self.atualizar()

    def carregar(self):
        cam = filedialog.askopenfilename(title="Planilha de constantes",
                                         filetypes=[("Excel", "*.xlsx *.xlsm"), ("Todos", "*.*")])
        if not cam:
            return
        try:
            itens = ler_planilha_constantes(cam)
        except Exception as ex:
            messagebox.showerror("Constantes", str(ex))
            return
        if not itens:
            messagebox.showinfo("Constantes", "Não há constantes na planilha.")
            return
        n = ligar_constantes(itens)
        if not messagebox.askyesno("Constantes", f"A planilha tem {len(itens)} constante(s); {n} foram ligadas "
                                                 "ao nome do material que já está no sistema" +
                                                 (f" e {len(itens) - n} não foram encontradas (ficam com o nome "
                                                  "da planilha, em laranja)" if n < len(itens) else "") +
                                                 f".\n\nSubstituir as {len(CONSTANTES.itens)} atuais?"):
            return
        CONSTANTES.recarregar_se_mudou()
        CONSTANTES.definir(itens)
        self.limpar()
        self.app.cadastro_mudou()
        self.atualizar()

    def baixar(self):
        cam = filedialog.asksaveasfilename(defaultextension=".xlsx", filetypes=[("Excel", "*.xlsx")],
                                           initialfile="constantes.xlsx")
        if not cam:
            return
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill
            wb = Workbook()
            ws = wb.active
            ws.append(["CÓDIGO", "DESCRIÇÃO", "CONSTANTE"])
            for c in ws[1]:
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = PatternFill("solid", fgColor="014FA0")
            for i in CONSTANTES.itens:
                cod = i.get("codigo", "")
                ws.append([int(cod) if cod.isdigit() else cod, i["descricao"], i["constante"]])
            for col, w in zip("ABC", (14, 45, 14)):
                ws.column_dimensions[col].width = w
            wb.save(cam)
        except Exception as ex:
            messagebox.showerror("Constantes", f"{ex}\n\nSe a planilha estiver aberta no Excel, feche e tente de novo.")
            return
        messagebox.showinfo("Constantes", f"Planilha salva:\n{cam}")


class TelaCarteira(_TelaBase):
    """Carteira de obra: tabela com os lançamentos, filtros, novo lançamento e edição."""

    def __init__(self, pai, app):
        super().__init__(pai, bg=FUNDO)
        self.app = app
        self.pode = True
        self.ordem = ("inicio", False)
        self._cabeca("Carteira de obra",
                     "Clique em \"+ Novo lançamento\" para lançar uma obra. Para alterar, dê dois cliques na "
                     "linha (ou selecione e clique em \"Editar\").")
        bar = cartao(self)
        bar.pack(fill="x", padx=20, pady=(0, 8))
        b = tk.Frame(bar, bg=BRANCO)
        b.pack(fill="x", padx=12, pady=(10, 4))
        ttk.Button(b, text="+ Novo lançamento", style="Verm.TButton", command=self.novo).pack(side="left")
        ttk.Button(b, text="Editar", command=self.editar).pack(side="left", padx=(6, 0))
        self._ed(ttk.Button(b, text="Excluir", command=self.excluir)).pack(side="left", padx=(4, 0))
        ttk.Separator(b, orient="vertical").pack(side="left", fill="y", padx=10)
        self._ed(ttk.Button(b, text="Importar planilha", command=self.importar)).pack(side="left")
        ttk.Button(b, text="Exportar Excel", style="Azul.TButton", command=self.exportar).pack(side="left", padx=4)
        self.lbl_n = tk.Label(b, text="", bg=BRANCO, fg=CINZA, font=(FONTE, 9))
        self.lbl_n.pack(side="right")
        ttk.Button(b, text="Limpar filtros", command=self.limpar_filtros).pack(side="right", padx=(0, 10))
        f = tk.Frame(bar, bg=BRANCO)
        f.pack(fill="x", padx=12, pady=(4, 10))
        tk.Label(f, text="Procurar:", bg=BRANCO, fg=CINZA, font=(FONTE, 9)).pack(side="left")
        self.v_busca = tk.StringVar()
        ttk.Entry(f, textvariable=self.v_busca, width=22).pack(side="left", padx=(4, 12))
        self.v_busca.trace_add("write", lambda *a: self.mostrar())
        self.v_f = {}
        for k, rot, w in (("status", "Status", 15), ("supervisor", "Supervisor", 13), ("encarregado", "Encarregado", 13),
                          ("equipe", "Equipe", 9)):
            tk.Label(f, text=rot + ":", bg=BRANCO, fg=CINZA, font=(FONTE, 9)).pack(side="left")
            v = tk.StringVar(value="Todos")
            cb = ttk.Combobox(f, textvariable=v, state="readonly", width=w)
            cb.pack(side="left", padx=(4, 10))
            cb.bind("<<ComboboxSelected>>", lambda e: self.mostrar())
            self.v_f[k] = (v, cb)
        tk.Label(f, text="Início — mês:", bg=BRANCO, fg=CINZA, font=(FONTE, 9)).pack(side="left")
        self.v_mes = tk.StringVar(value="Todos")
        self.cb_mes = ttk.Combobox(f, textvariable=self.v_mes, state="readonly", width=10,
                                   values=("Todos",) + MESES + ("Sem data",))
        self.cb_mes.pack(side="left", padx=(4, 6))
        self.cb_mes.bind("<<ComboboxSelected>>", lambda e: self.mostrar())
        tk.Label(f, text="ano:", bg=BRANCO, fg=CINZA, font=(FONTE, 9)).pack(side="left")
        self.v_ano = tk.StringVar(value="Todos")
        self.cb_ano = ttk.Combobox(f, textvariable=self.v_ano, state="readonly", width=6)
        self.cb_ano.pack(side="left", padx=(4, 0))
        self.cb_ano.bind("<<ComboboxSelected>>", lambda e: self.mostrar())
        self.resumo = tk.Frame(bar, bg=BRANCO)
        self.resumo.pack(fill="x", padx=12, pady=(0, 8))

        tab = cartao(self)
        tab.pack(fill="both", expand=True, padx=20, pady=(0, 16))
        fr = tk.Frame(tab, bg=BRANCO)
        fr.pack(fill="both", expand=True, padx=8, pady=8)
        cols = [k for k, _, _ in CAMPOS_OBRA]
        self.tree = ttk.Treeview(fr, columns=cols, show="headings", height=8)
        for k, t, w in CAMPOS_OBRA:
            self.tree.heading(k, text=t, command=lambda c=k: self.ordenar(c))
            self.tree.column(k, width=w, minwidth=50, stretch=False, anchor="w" if k == "obs" else "center")
        sy = ttk.Scrollbar(fr, orient="vertical", command=self.tree.yview)
        sx = ttk.Scrollbar(fr, orient="horizontal", command=self.tree.xview)
        self.tree.configure(yscrollcommand=sy.set, xscrollcommand=sx.set)
        self.tree.grid(row=0, column=0, sticky="nsew")
        sy.grid(row=0, column=1, sticky="ns")
        sx.grid(row=1, column=0, sticky="ew")
        fr.rowconfigure(0, weight=1)
        fr.columnconfigure(0, weight=1)
        for st, (fundo, _) in CORES_STATUS.items():
            self.tree.tag_configure(st, background=fundo)
        self.tree.bind("<Double-1>", lambda e: self.editar() if self.tree.identify_region(e.x, e.y) == "cell"
                       else None)
        self.tree.bind("<Return>", lambda e: self.editar())
        self.tree.bind("<Delete>", lambda e: self.excluir())

    # ---------------------------------------------------------------- tabela
    def atualizar(self):
        CARTEIRA.recarregar_se_mudou()
        EQUIPES.recarregar_se_mudou()
        for k, (v, cb) in self.v_f.items():
            if k == "status":
                vals = list(STATUS_OBRA)
            else:
                vals = sorted({o.get(k, "") for o in CARTEIRA.obras if o.get(k)} |
                              set(EQUIPES.codigos() if k == "equipe" else
                                  EQUIPES.nomes(k.capitalize())), key=str.upper)
            cb.configure(values=["Todos"] + vals)
        anos = {data_ordem(o.get("inicio"))[0] for o in CARTEIRA.obras} - {9999}
        anos.add(time.localtime().tm_year)
        self.cb_ano.configure(values=["Todos"] + [str(a) for a in sorted(anos, reverse=True)])
        self.mostrar()

    def limpar_filtros(self):
        self.v_busca.set("")
        self.v_mes.set("Todos")
        self.v_ano.set("Todos")
        for v, _ in self.v_f.values():
            v.set("Todos")
        self.mostrar()

    def _filtradas(self):
        k = chave_nome(self.v_busca.get())
        mes, ano = self.v_mes.get(), self.v_ano.get()
        lst = []
        for o in CARTEIRA.obras:
            a, m, _ = data_ordem(o.get("inicio"))           # filtro pela data de início
            if mes == "Sem data":
                if a != 9999:
                    continue
            elif mes != "Todos" and (a == 9999 or MESES[m - 1] != mes):
                continue
            if ano != "Todos" and str(a) != ano:
                continue
            if any(v.get() != "Todos" and o.get(c, "") != v.get() for c, (v, _) in self.v_f.items()):
                continue
            if k and k not in chave_nome(" ".join(str(o.get(c, "")) for c, _, _ in CAMPOS_OBRA)):
                continue
            lst.append(o)
        c, inv = self.ordem
        if c in ("inicio", "fim", "conclusao"):
            chave = lambda o: (data_ordem(o.get(c)), o.get("projeto", ""))
        elif c in ("projeto", "nota"):
            chave = lambda o: (len(str(o.get(c, ""))), str(o.get(c, "")))
        else:
            chave = lambda o: (str(o.get(c, "")).upper(), data_ordem(o.get("inicio")))
        return sorted(lst, key=chave, reverse=inv)

    def mostrar(self):
        sel = self.tree.selection()
        self.tree.delete(*self.tree.get_children())
        lst = self._filtradas()
        for o in lst:
            self.tree.insert("", "end", iid=str(o["id"]),
                             values=[o.get(k, "").replace("\n", " ") for k, _, _ in CAMPOS_OBRA],
                             tags=(o.get("status", ""),))
        for s in sel:
            if self.tree.exists(s):
                self.tree.selection_set(s)
                self.tree.see(s)
        c, inv = self.ordem
        for k, t, _ in CAMPOS_OBRA:
            self.tree.heading(k, text=t + ((" ▼" if inv else " ▲") if k == c else ""))
        self.lbl_n.config(text=f"{len(lst)} de {len(CARTEIRA.obras)} lançamento(s)")
        for w in self.resumo.winfo_children():
            w.destroy()
        for st in STATUS_OBRA:
            n = sum(1 for o in lst if o.get("status") == st)
            fundo, cor = CORES_STATUS[st]
            tk.Label(self.resumo, text=f" {st}: {n} ", bg=fundo, fg=cor, font=(FONTE, 9, "bold"),
                     padx=6, pady=2).pack(side="left", padx=(0, 6))

    def ordenar(self, c):
        self.ordem = (c, not self.ordem[1] if self.ordem[0] == c else False)
        self.mostrar()

    def _sel(self):
        s = self.tree.selection()
        if not s:
            messagebox.showinfo("Carteira de obra", "Selecione um lançamento na tabela.")
            return None
        return CARTEIRA.obra(int(s[0]))

    def _usuario(self):
        return self.app.usuario.get("nome") or self.app.usuario.get("login") or ""

    def novo(self):
        EQUIPES.recarregar_se_mudou()
        r = FormObra(self, None, self._usuario()).resultado
        if r is not None:
            self.atualizar()
            self.tree.selection_set(str(r))
            self.tree.see(str(r))
            self.app.msg("Lançamento incluído na carteira.")

    def editar(self):
        CARTEIRA.recarregar_se_mudou()
        o = self._sel()
        if not o:
            return
        r = FormObra(self, dict(o), self._usuario()).resultado
        if r is not None:
            self.atualizar()
            self.tree.selection_set(str(r))
            self.tree.see(str(r))
            self.app.msg("Lançamento alterado.")

    def excluir(self):
        if not self.pode:
            return
        s = self.tree.selection()
        if not s:
            messagebox.showinfo("Carteira de obra", "Selecione na tabela o(s) lançamento(s) que quer excluir.")
            return
        if not messagebox.askyesno("Carteira de obra", f"Excluir {len(s)} lançamento(s)?\n\nIsso não pode ser "
                                                       "desfeito."):
            return
        for i in s:
            CARTEIRA.excluir_obra(int(i))
        self.atualizar()

    def exportar(self):
        lst = self._filtradas()
        if not lst:
            messagebox.showinfo("Exportar", "Não há lançamentos para exportar.")
            return
        cam = filedialog.asksaveasfilename(defaultextension=".xlsx", filetypes=[("Excel", "*.xlsx")],
                                           initialfile=f"carteira_de_obra_{time.strftime('%d-%m-%Y')}.xlsx")
        if not cam:
            return
        try:
            salvar_planilha_carteira(cam, lst)
        except Exception as ex:
            messagebox.showerror("Exportar", f"{ex}\n\nSe a planilha estiver aberta no Excel, feche e tente de novo.")
            return
        messagebox.showinfo("Exportar", f"{len(lst)} lançamento(s) salvos em:\n{cam}"
                                        "\n\n(Sai o que está aparecendo na tabela, com os filtros.)")

    def importar(self):
        if not self.pode:
            return
        cam = filedialog.askopenfilename(title="Planilha da carteira de obra",
                                         filetypes=[("Excel", "*.xlsx *.xlsm"), ("Todos", "*.*")])
        if not cam:
            return
        try:
            novas = ler_planilha_carteira(cam)
        except Exception as ex:
            messagebox.showerror("Importar", str(ex))
            return
        if not novas:
            messagebox.showinfo("Importar", "Não há lançamentos na planilha.")
            return
        CARTEIRA.recarregar_se_mudou()
        dups = [d for d in novas if CARTEIRA.duplicada(d["projeto"], d.get("nota", ""))]
        atualizar = False
        if dups:
            r = messagebox.askyesnocancel("Importar", f"A planilha tem {len(novas)} lançamento(s); {len(dups)} já "
                                                      "estão na carteira (mesmo projeto e nota).\n\n"
                                                      "Sim = atualizar esses com os dados da planilha\n"
                                                      "Não = deixar esses como estão e trazer só os novos")
            if r is None:
                return
            atualizar = r
        # pessoas e equipes que não estão cadastradas
        faltam = {}
        for d in novas:
            for k, fn in (("supervisor", "Supervisor"), ("encarregado", "Encarregado")):
                if d.get(k) and not EQUIPES.pessoa(d[k]):
                    faltam.setdefault(("p", d[k]), fn)
            if d.get("equipe") and not EQUIPES.equipe(d["equipe"]):
                faltam.setdefault(("e", d["equipe"]), (d.get("supervisor", ""), d.get("encarregado", "")))
        cadastrar = False
        if faltam:
            np_ = sum(1 for k in faltam if k[0] == "p")
            ne = len(faltam) - np_
            cadastrar = messagebox.askyesno("Importar", f"A planilha tem {np_} pessoa(s) e {ne} equipe(s) que "
                                                        "ainda não estão no Cadastro de equipe.\n\n"
                                                        "Cadastrar automaticamente?")
        n_novo = n_atual = 0
        for d in novas:
            dup = CARTEIRA.duplicada(d["projeto"], d.get("nota", ""))
            if dup:
                if atualizar:
                    CARTEIRA.salvar_obra({**dup, **d}, dup["id"], self._usuario())
                    n_atual += 1
                continue
            CARTEIRA.salvar_obra(d, None, self._usuario())
            n_novo += 1
        if cadastrar:
            for (tipo, nome), extra in faltam.items():
                try:
                    if tipo == "p":
                        EQUIPES.salvar_pessoa(None, nome, extra)
                    else:
                        EQUIPES.salvar_equipe(None, nome, extra[0], extra[1], [])
                except ValueError:
                    pass
        self.atualizar()
        messagebox.showinfo("Importar", f"{n_novo} lançamento(s) novos" +
                            (f", {n_atual} atualizado(s)" if n_atual else "") + "." +
                            ("\n\nPessoas e equipes cadastradas: confira os colaboradores de cada equipe em "
                             "Cadastro de equipe." if cadastrar else ""))

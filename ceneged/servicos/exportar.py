"""Exportação do relatório para Excel (NEO, conferência, todos os projetos).

Parte do programa Levantamento de Materiais (Ceneged)."""

import csv
import re

from ceneged.comum.texto import chave_nome
from ceneged.dados.composicoes import MOVIMENTOS

COLUNAS_NEO = ["Codigo_Material", "Material", "Quantidade", "Serial", "Valor", "Situacao", "NCA",
               "Fabricante", "Patrimonio", "Caixa", "Local"]


def _aba_neo(ws, linhas):
    """Aba no formato da Neoenergia: Codigo_Material | Material | Quantidade | ... | Situacao=NV.
    Entram todos os materiais; os que não têm código ficam com a linha em vermelho."""
    from openpyxl.styles import Font, PatternFill
    ws.append(COLUNAS_NEO)
    verm = PatternFill("solid", fgColor="FDE4E4")
    for lin in linhas:
        cod = lin.get("codigo") or ""
        q = round(float(lin["qtd"]), 3)
        ws.append([int(cod) if cod.isdigit() else (cod or None), lin["material"],
                   int(q) if q == int(q) else q, None, None, "NV", None, None, None, None, None])
        if not cod:
            for c in ws[ws.max_row]:
                c.fill = verm
                c.font = Font(color="BB1B21")
    for col, w in zip("ABCDEFGHIJK", (16, 46, 12, 10, 10, 10, 10, 12, 12, 10, 10)):
        ws.column_dimensions[col].width = w


def _aba_conferencia(ws, projetos):
    """Conferência: projeto, código, descrição, unidade, quantidade, de onde veio e pendência.
    projetos: [(nome do projeto, linhas)]"""
    from openpyxl.styles import Font, PatternFill
    ws.append(["Projeto", "Codigo_Material", "Descrição", "Und", "Quantidade", "Movimento", "Veio de",
               "Pendência"])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="014FA0")
    verm = Font(color="BB1B21")
    for nome, linhas in projetos:
        for lin in linhas:
            cod = lin.get("codigo") or ""
            q = round(float(lin["qtd"]), 3)
            orig = []
            if "projeto" in lin["origens"] or "cabo" in lin["origens"]:
                orig.append("balão")
            if lin.get("de"):
                orig.append("material principal " + lin["de"])
            if "manual" in lin["origens"]:
                orig.append("ajuste manual")
            ws.append([nome, int(cod) if cod.isdigit() else cod, lin["material"], lin["unidade"],
                       int(q) if q == int(q) else q, lin.get("mov") or MOVIMENTOS[0], ", ".join(orig),
                       lin.get("pendente") or ("" if cod else "sem código")])
            if not cod:
                for c in ws[ws.max_row]:
                    c.font = verm
    for col, w in zip("ABCDEFGH", (22, 16, 48, 6, 12, 20, 36, 50)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:H{max(ws.max_row, 1)}"


def _total_dos_projetos(projetos):
    """Soma de todos os projetos por código (ou nome, sem código) e unidade."""
    tot = {}
    for _, linhas in projetos:
        for lin in linhas:
            k = (lin.get("codigo") or "n:" + chave_nome(lin["material"]), lin["unidade"])
            if k not in tot:
                tot[k] = dict(lin, qtd=0.0)
            tot[k]["qtd"] += float(lin["qtd"])
    return sorted(tot.values(), key=lambda l: (not l.get("codigo"), l["material"]))


def _nome_aba(nome, usados):
    n = re.sub(r"[\[\]:*?/\\]", "_", nome).strip()[:31] or "Projeto"
    base, i = n, 2
    while n.upper() in usados:
        suf = f" ({i})"
        n = base[:31 - len(suf)] + suf
        i += 1
    usados.add(n.upper())
    return n


def salvar_relatorio_neo(caminho, linhas, conferencia_de=None):
    """Um projeto: aba "Planilha1" no formato da Neoenergia (e, se pedido, a aba Conferência
    no mesmo arquivo)."""
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Planilha1"
    _aba_neo(ws, linhas)
    if conferencia_de:
        _aba_conferencia(wb.create_sheet("Conferência"), [(conferencia_de, linhas)])
    wb.save(caminho)


def salvar_relatorio_todos(caminho, projetos, conferencia=True):
    """Todos os projetos num arquivo só: uma aba por projeto (formato da Neoenergia), a aba
    TOTAL (soma de todos) e, se pedido, a aba Conferência com todos os projetos."""
    from openpyxl import Workbook
    wb = Workbook()
    usados = set()
    primeiro = True
    for nome, linhas in projetos:
        ws = wb.active if primeiro else wb.create_sheet()
        ws.title = _nome_aba(nome, usados)
        _aba_neo(ws, linhas)
        primeiro = False
    if len(projetos) > 1:
        _aba_neo(wb.create_sheet(_nome_aba("TOTAL", usados)), _total_dos_projetos(projetos))
    if conferencia:
        _aba_conferencia(wb.create_sheet(_nome_aba("Conferência", usados)), projetos)
    wb.save(caminho)


def salvar_conferencia(caminho, nome, linhas):
    """Planilha de conferência separada (mantida para compatibilidade)."""
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Conferência"
    _aba_conferencia(ws, [(nome, linhas)])
    wb.save(caminho)


def salvar_relatorio(caminho, linhas):
    """linhas: [(arquivo, material, unidade, qtd)]. Salva .xlsx (ou .csv)."""
    if caminho.lower().endswith(".xlsx"):
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        wb = Workbook()
        ws = wb.active
        ws.title = "Relatório"
        ws.append(["Arquivo", "Material", "Quantidade"])
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="305496")
        for arq, mat, un, q in linhas:
            ws.append([arq, mat, q])
            if un == "M":
                ws.cell(ws.max_row, 3).number_format = '0" M"'
        larg = [max(len(str(r[i])) for r in [("Arquivo", "Material", "Quantidade")] +
                    [(a, m, f"{q} M") for a, m, _, q in linhas]) for i in range(3)]
        for col, w in zip("ABC", larg):
            ws.column_dimensions[col].width = min(max(w + 2, 12), 70)
        ws.freeze_panes = "A2"
        wb.save(caminho)
    else:
        with open(caminho, "w", newline="", encoding="utf-8-sig") as f:
            w = csv.writer(f, delimiter=";")
            w.writerow(["Arquivo", "Material", "Quantidade"])
            for arq, mat, un, q in linhas:
                w.writerow([arq, mat, f"{q} M" if un == "M" else q])
